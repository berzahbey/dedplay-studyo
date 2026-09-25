"""Kitap Okuma'nın temiz metin parçalarından (Türkçe) ve Osmanlıca karşılıklarından kitap dosyaları üretir.

variant: "tr" (Türkçe), "osm" (Osmanlıca), "iki" (Türkçe + Osmanlıca)
"""
import html
import io
import os
import re
import tempfile

AMIRI = "/usr/share/fonts/truetype/amiri/Amiri-Regular.ttf"
CHUNK = 20  # EPUB'da bir bölüm dosyasına kaç parça girsin

LABEL = {"tr": "Türkçe", "osm": "Osmanlıca", "iki": "Türkçe ve Osmanlıca"}


def _paras(text):
    return [p.strip() for p in re.split(r"\n+", text or "") if p.strip()]


def _blocks(parts, variant):
    """Her parça için [(tür, paragraf), ...] listesi."""
    out = []
    for p in parts:
        b = []
        if variant == "iki":
            tr, osm = _paras(p["tr"]), _paras(p["osm"])
            if len(tr) == len(osm):
                for a, o in zip(tr, osm):
                    b += [("tr", a), ("osm", o)]
            else:
                b += [("tr", a) for a in tr] + [("osm", o) for o in osm]
        elif variant == "osm":
            b = [("osm", o) for o in _paras(p["osm"])]
        else:
            b = [("tr", a) for a in _paras(p["tr"])]
        out.append(b)
    return out


def _title(job, variant):
    if variant == "osm" and job["osm_title"]:
        return job["osm_title"]
    return job["title"]


def _p(kind, text):
    t = html.escape(text)
    if kind == "osm":
        return f'<p class="osm" dir="rtl" lang="ota">{t}</p>'
    return f'<p class="tr">{t}</p>'


# ---------------- TXT ----------------
def to_txt(job, parts, variant):
    lines = [_title(job, variant), ""]
    for b in _blocks(parts, variant):
        for _, t in b:
            lines += [t, ""]
    return "\n".join(lines).encode("utf-8")


# ---------------- HTML / PDF ----------------
BASE_CSS = """
body { font-family: "Noto Serif", Georgia, serif; line-height: 1.7; color: #1c1c1c; }
p { margin: 0 0 .9em; text-align: justify; }
p.osm { font-family: "Amiri", "Noto Naskh Arabic", serif; font-size: 1.3em; line-height: 1.9; text-align: justify; direction: rtl; }
.iki p.tr { margin-bottom: .2em; color: #444; }
.iki p.osm { margin-bottom: 1.2em; }
.title { text-align: center; }
body.osm .title h1 { font-family: "Amiri", "Noto Naskh Arabic", serif; }
"""
HTML_CSS = BASE_CSS + """
body { max-width: 40em; margin: 0 auto; padding: 2em 1.2em 4em; background: #fdfdfb; }
.title { margin: 3em 0 4em; } .title h1 { font-size: 2em; line-height: 1.3; } .title p { color: #777; }
"""
PDF_CSS = BASE_CSS + """
@page { size: A5; margin: 18mm 15mm 20mm;
        @bottom-center { content: counter(page); font-family: "Noto Serif", serif; font-size: 8.5pt; color: #888; } }
@page :first { @bottom-center { content: none; } }
body { font-size: 10pt; line-height: 1.55; }
p { margin-bottom: 7pt; orphans: 2; widows: 2; }
p.osm { font-size: 13pt; line-height: 1.8; }
.title { page-break-after: always; padding-top: 35%; }
.title h1 { font-size: 20pt; line-height: 1.4; } .title p { color: #777; font-size: 10pt; }
"""
FONT_LINK = ('<link rel="stylesheet" href="https://fonts.googleapis.com/css2?'
             'family=Amiri&family=Noto+Serif:wght@400;700&display=swap">')


def to_html(job, parts, variant, css=HTML_CSS, head=FONT_LINK):
    title = html.escape(_title(job, variant))
    tdir = ' dir="rtl"' if variant == "osm" else ""
    body = "\n".join(_p(k, t) for b in _blocks(parts, variant) for k, t in b)
    lang = "ota" if variant == "osm" else "tr"
    doc = (f'<!doctype html><html lang="{lang}"><head><meta charset="utf-8">'
           f'<meta name="viewport" content="width=device-width, initial-scale=1"><title>{title}</title>'
           f'{head}<style>{css}</style></head><body class="{variant}">'
           f'<div class="title"><h1{tdir}>{title}</h1><p>{LABEL[variant]}</p></div>{body}</body></html>')
    return doc.encode("utf-8")


def to_pdf(job, parts, variant):
    from weasyprint import HTML
    return HTML(string=to_html(job, parts, variant, css=PDF_CSS, head="").decode("utf-8")).write_pdf()


# ---------------- EPUB ----------------
def to_epub(job, parts, variant):
    from ebooklib import epub

    title = _title(job, variant)
    book = epub.EpubBook()
    book.set_identifier(f"dedplay-studyo-{job['id']}-{variant}")
    book.set_title(title)
    book.set_language("ota" if variant == "osm" else "tr")
    if variant == "osm":
        book.set_direction("rtl")
    css = BASE_CSS
    if variant != "tr" and os.path.exists(AMIRI):
        book.add_item(epub.EpubItem(uid="amiri", file_name="fonts/Amiri.ttf", media_type="font/ttf",
                                    content=open(AMIRI, "rb").read()))
        css = '@font-face { font-family: "Amiri"; src: url(fonts/Amiri.ttf); }\n' + css
    style = epub.EpubItem(uid="style", file_name="style.css", media_type="text/css", content=css)
    book.add_item(style)

    blocks = _blocks(parts, variant)
    chapters = []
    for n, start in enumerate(range(0, max(1, len(blocks)), CHUNK), 1):
        chunk = blocks[start:start + CHUNK]
        name = f"Kısım {n}"
        body = "\n".join(_p(k, t) for b in chunk for k, t in b)
        ch = epub.EpubHtml(title=name, file_name=f"kisim_{n:04d}.xhtml", lang="ota" if variant == "osm" else "tr")
        ch.content = f'<html><body class="{variant}">{body}</body></html>'
        ch.add_item(style)
        book.add_item(ch)
        chapters.append(ch)
    book.toc = chapters
    book.add_item(epub.EpubNcx())
    book.add_item(epub.EpubNav())
    book.spine = ["nav"] + chapters
    with tempfile.NamedTemporaryFile(suffix=".epub", delete=False) as f:
        tmp = f.name
    try:
        epub.write_epub(tmp, book)
        return open(tmp, "rb").read()
    finally:
        os.remove(tmp)


# ---------------- DOCX ----------------
def _rtl(paragraph):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    paragraph._p.get_or_add_pPr().append(OxmlElement("w:bidi"))
    for run in paragraph.runs:
        rpr = run._r.get_or_add_rPr()
        rpr.append(OxmlElement("w:rtl"))
        fonts = rpr.find(qn("w:rFonts"))
        if fonts is None:
            fonts = OxmlElement("w:rFonts")
            rpr.insert(0, fonts)
        fonts.set(qn("w:cs"), "Amiri")


def to_docx(job, parts, variant):
    import docx
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt, RGBColor

    d = docx.Document()
    d.core_properties.title = _title(job, variant)
    h = d.add_heading(_title(job, variant), level=0)
    if variant == "osm":
        _rtl(h)
        h.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for b in _blocks(parts, variant):
        for kind, t in b:
            p = d.add_paragraph()
            r = p.add_run(t)
            if kind == "osm":
                r.font.size = Pt(15)
                _rtl(p)
                p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            elif variant == "iki":
                r.font.color.rgb = RGBColor(0x55, 0x55, 0x55)
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


FORMATS = {
    "epub": (to_epub, "application/epub+zip"),
    "pdf": (to_pdf, "application/pdf"),
    "docx": (to_docx, "application/vnd.openxmlformats-officedocument.wordprocessingml.document"),
    "html": (to_html, "text/html; charset=utf-8"),
    "txt": (to_txt, "text/plain; charset=utf-8"),
}
SUFFIX = {"tr": "Türkçe", "osm": "Osmanlıca", "iki": "Türkçe-Osmanlıca"}


def build(job, parts, fmt, variant):
    fn, media = FORMATS[fmt]
    base = f"{job['title']} - {SUFFIX[variant]}"
    if job["status"] != "done":
        base += " (kısmi)"
    return fn(job, parts, variant), media, f"{base}.{fmt}"
