"""Elindeki kitaplar (PDF, EPUB, DOCX, TXT) -> kitap.json (Türkçe).

Stüdyo'nun metin kodu (app.textsrc, app.duzelt) kullanılır; ondan farklı olarak burada
basılı sayfa numaraları, başlıklar (fihrist için) ve dipnot bağlantıları KORUNUR.

İç işaretler (özel kullanım alanı karakterleri):
  \\ue000n\\ue001  metindeki dipnot numarası (sayfanın dipnotuna bağlanmadan önce)
  \\ue002etiket\\ue003  basılı sayfa başlangıcı (düzeltmeden sonra konuma çevrilir)
  {{n0001}}  bağlanmış dipnot atfı (kitap.json biçimi)
"""
import collections
import hashlib
import io
import os
import html
import re
import statistics as st
import unicodedata
from concurrent.futures import ProcessPoolExecutor

from . import kitap as K
from . import zeyrek_onarim  # noqa: F401  zeyrek'in kök kümesini bozan hatası (DZ'den önce yüklenmeli)

try:  # Stüdyo'nun kodu (imajda /app/app)
    from app import duzelt as DZ
    from app import textsrc as TS
except ImportError:  # pragma: no cover
    DZ = TS = None

if DZ is not None and not getattr(DZ.gecerli_mi, "_kilitli", False):
    # 0.5.17: iki iş şeridi Zemberek'i aynı anda kullanabilir; çözümleyici iş parçacığı güvenli değil. Bütün çözümleme
    # DZ.gecerli_mi'den geçer (Stüdyo'nun kendi fonksiyonları da onu modül adıyla çağırır)
    import threading as _th
    _zemberek_kilidi = _th.Lock()
    _asil_gecerli_mi = DZ.gecerli_mi

    def _kilitli_gecerli_mi(w):
        with _zemberek_kilidi:
            return _asil_gecerli_mi(w)
    _kilitli_gecerli_mi._kilitli = True
    DZ.gecerli_mi = _kilitli_gecerli_mi

UST = re.compile(r"\ue000(\d{1,3})\ue001")
SAYFA_ISARET = re.compile(r"\ue002([^\ue003]{1,20})\ue003")
END_PUNCT = tuple('.!?:;"”»)]…')
YAPISIK = re.compile(r"(?<=[a-zçğıöşüâîû\.,;:\"”»’')])(\d{1,2})(?=[\s.,;:!?”\"')]|$)")  # 0.5.16: tırnaktan sonra
BOLUM_NO = re.compile(r"^(birinci|ikinci|üçüncü|dördüncü|beşinci|altıncı|yedinci|sekizinci|dokuzuncu|onuncu|"
                      r"\d{1,2}\.?|[ivxlc]{1,6}\.?)\s+(bölüm|kısım|fasıl|bab|kitap|makale|mektup|risale)[.:]?$", re.I)
OCR_DPI = 300


def _gerekli():
    if TS is None or DZ is None:
        raise RuntimeError("Stüdyo'nun metin kodu (app.textsrc, app.duzelt) bulunamadı")


def _harf(s):
    return sum(1 for c in s if c.isalpha())


def _rakam(s):
    return s.translate(str.maketrans("٠١٢٣٤٥٦٧٨٩۰۱۲۳۴۵۶۷۸۹", "01234567890123456789"))


# ======================= PDF: satırlar =======================
_TR_PARCA = re.compile(r"\s*[a-zA-Zçğıöşüâîû]{0,2}[ıİğĞşŞçÇöÖüÜ][a-zA-Zçğıöşüâîû]{0,2}\s*")


def _katman_satirlari(page):
    rows = []
    bloklar = page.get_text("dict").get("blocks", [])
    # metin sütununun sağ kenarı (uzun satırların ortancası): bunun dışındaki "(17)" parçaları kenar numarasıdır
    uzun = sorted(ln["bbox"][2] for b in bloklar for ln in b.get("lines", [])
                  if len("".join(s.get("text", "") for s in ln.get("spans", []))) > 40)
    sag_kenar = uzun[len(uzun) // 2] if len(uzun) >= 5 else None
    for b in bloklar:
        for ln in b.get("lines", []):
            spans = [s for s in ln.get("spans", []) if s.get("text", "").strip()]
            if sag_kenar:
                spans = [s for s in spans if not (s["bbox"][0] >= sag_kenar + 2 and
                                                  re.fullmatch(r"\s*\(?\s*\d{1,4}\s*\)?\s*", s["text"]))]
            if not spans:
                continue
            boy = [s["size"] for s in spans if any(c.isalpha() for c in s["text"])]
            h = st.median(boy) if boy else (spans[0]["size"] if spans else 0)
            parca = []
            for k, s in enumerate(spans):
                t = s["text"]
                if k > 0 and _TR_PARCA.fullmatch(t) and h and s["size"] < h * 0.92:
                    parca.append(t.lstrip())  # OCR katmanı: ı/ğ/ş ayrı yazı tipinde, önüne boşluk konmuş
                    continue
                ust = (s.get("flags", 0) & 1) or (h and s["size"] < h * 0.78)
                if k > 0 and ust and re.fullmatch(r"\s*[\d٠-٩]{1,3}\s*", t):
                    parca.append("\ue000" + _rakam(t.strip()) + "\ue001")
                    continue
                if k > 0 and ust and re.fullmatch(r"[\d\s\W]+", t):
                    continue
                parca.append(t)
            metin = TS.norm("".join(parca))
            if not metin:
                continue
            harfli = [s for s in spans if any(c.isalpha() for c in s["text"])]
            kalin = bool(harfli) and all(("bold" in s.get("font", "").lower() or s.get("flags", 0) & 16) for s in harfli)
            x0, y0, x1, y1 = ln["bbox"]
            rows.append({"text": metin, "h": h, "top": y0, "bot": y1, "x0": x0, "x1": x1,
                         "n": len(metin.split()), "kalin": kalin, "blok": b.get("number", 0), "ocr": False})
    return _satir_parcalarini_birlestir(rows)


def _satir_parcalarini_birlestir(rows):
    """Eski OCR katmanı bir satırı aynı yükseklikte iki-üç ayrı satır olarak kaydetmiş olabilir ("…toplumdaki yeri,
    onların" + "İslâm toplumundaki yerini"); ayrı kalırsa paragraf cümle ortasında bölünür. Birleşme şartı: dikeyde en
    az %60 örtüşme, yakın yükseklik (dipnot işareti / içindekiler numarası karışmasın), aradaki boşluk en çok 2,5 harf
    yüksekliği (çift sayfa arası ve içindekilerin sağdaki numarası birleşmesin)."""
    if len(rows) < 2:
        return rows
    out = []
    for r in sorted(rows, key=lambda r: (round((r["top"] + r["bot"]) / 2), r["x0"])):
        es = None
        for o in reversed(out[-6:]):
            ust, alt = max(o["top"], r["top"]), min(o["bot"], r["bot"])
            yuk = min(o["bot"] - o["top"], r["bot"] - r["top"])
            hmax, hmin = max(o["h"], r["h"]), min(o["h"], r["h"]) or 1
            if yuk > 0 and alt - ust >= 0.6 * yuk and hmax / hmin <= 1.33 and (
                    -2 <= r["x0"] - o["x1"] <= 2.5 * hmax or -2 <= o["x0"] - r["x1"] <= 2.5 * hmax):
                es = o
                break
        if es is None:
            out.append(dict(r))
            continue
        sol, sag = (es, r) if es["x0"] <= r["x0"] else (r, es)  # parça sıralamada önce gelse de solda olabilir
        es["text"] = TS.norm(sol["text"] + " " + sag["text"])
        es["x0"] = min(es["x0"], r["x0"])
        es["x1"], es["top"], es["bot"] = max(es["x1"], r["x1"]), min(es["top"], r["top"]), max(es["bot"], r["bot"])
        es["n"] = len(es["text"].split())
        es["kalin"] = es["kalin"] and r["kalin"]
    if len(out) == len(rows):
        return rows
    # birleşme olduysa bir geçiş daha: büyüyen satır, daha önce geçilmiş parçayla ("bir" + "köprüdür, onun") birleşebilir
    return _satir_parcalarini_birlestir(sorted(out, key=lambda r: (r["blok"], r["top"])))


def _ocr_sayfa(args):
    """Taranmış sayfa: (i, satırlar, genişlik, yükseklik). Satır ölçüleri PDF birimine çevrilir."""
    os.environ["OMP_THREAD_LIMIT"] = "1"
    import fitz
    import pytesseract
    from PIL import Image, ImageOps
    yol, i, dil = args[:3]
    ayar = args[3] if len(args) > 3 else ""
    page = fitz.open(yol)[i]
    olcek = 72 / OCR_DPI
    pix = page.get_pixmap(dpi=OCR_DPI)
    img = ImageOps.autocontrast(Image.open(io.BytesIO(pix.tobytes("png"))).convert("L"))
    try:
        d = pytesseract.image_to_data(img, lang=dil, config=ayar, output_type=pytesseract.Output.DICT)
    except Exception:
        d = pytesseract.image_to_data(img, config=ayar, output_type=pytesseract.Output.DICT)
    if not ayar and dil != "ara":
        _arapca_yeniden_oku(img, d)
    rows = _ocr_satirlari(d, olcek, ayar)
    rows += _kenar_numarasi(img, rows, olcek, dil)
    return i, rows, page.rect.width, page.rect.height


_ARAP_HARF = re.compile(r"[\u0621-\u064A\u0671-\u06D3]")
_PARANTEZ = {"(", ")", "[", "]", "©", "®", "«", "»", "((", "))"}
_ara_var = None


def _ara_modeli_var():
    global _ara_var
    if _ara_var is None:
        try:
            import pytesseract
            _ara_var = "ara" in pytesseract.get_languages(config="")
        except Exception:
            _ara_var = False
    return _ara_var


def _arapca_yeniden_oku(img, d):
    """Türkçe kitaptaki Arapça ibare (ayet, hadis; paragraf içinde parantezle de) Türkçe modelle çöpe döner
    ("( © Öte op» Yİ yas )"). Satırda yan yana duran zayıf kelimeler (güven < 65, harfsiz ya da en çok 2 harfli; en az
    2 kelime, en az biri güven < 60) kesilip Arapça modelle yeniden okunur. Kabul şartı: kelimelerin çoğu Arapça harfli,
    en az 4 Arapça harf, Arapça kelimelerin güven ortalaması ≥ 45 ve en az birinin güveni ≥ 80. Kabul edilmezse satır
    olduğu gibi kalır (soluk basılmış Türkçe kelime Arapça sanılmasın). Kenardaki parantez korunur. d yerinde değişir."""
    if not _ara_modeli_var():
        return 0
    import pytesseract
    n = len(d["text"])
    satirlar = {}
    for k in range(n):
        if (d["text"][k] or "").strip():
            satirlar.setdefault((d["block_num"][k], d["par_num"][k], d["line_num"][k]), []).append(k)

    def zayif(k):
        t, g = d["text"][k].strip(), float(d["conf"][k])
        harf = sum(c.isalpha() for c in t)
        return g < 65 or harf == 0 or (harf <= 2 and not t.isdigit())

    degisen = 0
    for ks in satirlar.values():
        ks.sort(key=lambda k: d["left"][k])
        parcalar, cur = [], []
        for k in ks:
            if zayif(k):
                cur.append(k)
            else:
                if cur:
                    parcalar.append(cur)
                cur = []
        if cur:
            parcalar.append(cur)
        for par in parcalar:
            ic = [k for k in par if d["text"][k].strip() not in _PARANTEZ]  # kenardaki parantezler kesite girmez
            while ic and d["text"][ic[0]].strip() in _PARANTEZ:
                ic.pop(0)
            if len(ic) < 2 or not any(float(d["conf"][k]) < 60 for k in ic):
                continue
            # kesit yalnız Arapça parçanın kendi kelimelerine göre (satırın en yüksek kutusuna göre değil): boşluk
            # büyük olursa üst/alt satır kesite girer ve okuma bozulur
            h = st.median(d["height"][k] for k in ic)
            x0, x1 = min(d["left"][k] for k in ic), max(d["left"][k] + d["width"][k] for k in ic)
            y0, y1 = min(d["top"][k] for k in ic), max(d["top"][k] + d["height"][k] for k in ic)
            sy0, sy1 = min(d["top"][k] for k in ks), max(d["top"][k] + d["height"][k] for k in ks)
            sh = max(d["height"][k] for k in ks)
            en_iyi = None
            for kx0, ky0, kx1, ky1 in ((x0 - h // 3, y0 - h // 3, x1 + h // 3, y1 + h // 3),  # dar: parçanın kendisi
                                       (x0 - h * 0.7, y0 - h * 0.7, x1 + h * 0.7, y1 + h * 0.7),
                                       (x0 - sh // 2, sy0 - sh // 2, x1 + sh // 2, sy1 + sh // 2)):  # satır boyu
                # en iyi kesit satırdan satıra değişir (harekeler, üst/alt satıra yakınlık): üçü denenir, en güvenlisi
                kesit = img.crop((max(0, int(kx0)), max(0, int(ky0)), int(kx1), int(ky1)))
                try:
                    a = pytesseract.image_to_data(kesit, lang="ara", config="--psm 7",
                                                  output_type=pytesseract.Output.DICT)
                except Exception:
                    continue
                kel = [(t.strip(), float(g)) for t, g in zip(a["text"], a["conf"]) if (t or "").strip()]
                arap = [(t, g) for t, g in kel if _ARAP_HARF.search(t)]
                harf = sum(len(_ARAP_HARF.findall(t)) for t, _ in arap)
                if not arap or len(arap) < 0.6 * len(kel) or harf < 4 or \
                        sum(g for _, g in arap) / len(arap) < 45 or max(g for _, g in arap) < 80:
                    continue
                ort = sum(g for _, g in arap) / len(arap)
                if en_iyi is None or ort > en_iyi[0]:
                    en_iyi = (ort, arap)
            if en_iyi is None:
                continue
            arap = en_iyi[1]
            ayet = " ".join(t for t, _ in arap)
            parantez = any(d["text"][k].strip() in _PARANTEZ for k in par) or \
                any(d["text"][k].strip() in _PARANTEZ for k in ks if k not in par and abs(ks.index(k) - ks.index(par[0])) <= 1)
            ilk = par[0]
            d["text"][ilk] = ("( " + ayet + " )") if parantez else ayet
            d["conf"][ilk] = 99
            d["left"][ilk] = min(d["left"][k] for k in par)
            d["width"][ilk] = max(d["left"][k] + d["width"][k] for k in par) - d["left"][ilk]
            for k in par[1:]:
                d["text"][k] = ""
            for k in ks:  # parçanın hemen dışında kalan tek başına parantez artık metinde
                if k not in par and d["text"][k].strip() in _PARANTEZ and abs(ks.index(k) - ks.index(par[0])) <= len(par):
                    if parantez:
                        d["text"][k] = ""
            degisen += 1
    return degisen


def _ocr_satirlari(d, olcek, ayar=""):
    """Tesseract kelimeleri -> satırlar (PDF birimiyle). Satır içinde 3 harf yüksekliğinden büyük boşlukta satır
    bölünür: Tesseract 5.5 yan yana iki kitap sayfasının (ya da iki sütunun) aynı yükseklikteki satırlarını tek satır
    olarak verebiliyor ("ÖNSÖZ tehlikelidir…"). Normal kelime arası bir harf genişliği kadardır."""
    satir, sira = {}, []
    for k in range(len(d["text"])):
        t = (d["text"][k] or "").strip()
        if not t:
            continue
        key = (d["block_num"][k], d["par_num"][k], d["line_num"][k])
        if key not in satir:
            satir[key] = []
            sira.append(key)
        satir[key].append((d["left"][k], t, d["height"][k], d["top"][k], d["width"][k]))
    rows = []
    for key in sira:
        tum = sorted(satir[key])
        hs = [h for _, t, h, _, _ in tum if any(c.isalpha() for c in t)]
        lh = st.median(hs) if hs else 0
        gruplar = [[tum[0]]]
        for w in tum[1:]:
            onceki = gruplar[-1][-1]
            if lh and w[0] - (onceki[0] + onceki[4]) > 3 * lh:
                gruplar.append([w])
            else:
                gruplar[-1].append(w)
        for ws in gruplar:
            parca = []
            for k, (_, t, h, _, _) in enumerate(ws):
                if k > 0 and lh and h < lh * 0.62 and re.fullmatch(r"[\d]{1,3}", t):
                    parca.append("\ue000" + t + "\ue001")  # küçük (üst simge) dipnot numarası
                elif k > 0 and lh and h < lh * 0.6 and re.fullmatch(r"[\d\W]+", t) and not ayar:
                    continue  # (içindekiler okumasında nokta dizisi korunur)
                else:
                    parca.append((" " if parca else "") + t)
            metin = TS.norm(re.sub(r"(?:^|\s)[|¦]+(?=\s|$)", "", "".join(parca)).strip())  # "|": kalem çizgisi
            if metin:
                rows.append({"text": metin, "h": lh * olcek, "top": min(w[3] for w in ws) * olcek,
                             "bot": max(w[3] + w[2] for w in ws) * olcek, "x0": ws[0][0] * olcek,
                             "x1": max(w[0] + w[4] for w in ws) * olcek, "n": len(metin.split()), "kalin": False,
                             "blok": key[0] * 1000 + key[1], "ocr": True})
    return _satir_parcalarini_birlestir(rows)  # Tesseract'ın parçalı verdiği aynı satır (boşluk ≤ 2,5 harf) birleşir


def _kenar_numarasi(img, rows, olcek, dil):
    """Tesseract tek başına duran sayfa numarasını çoğu zaman görmez: üst ve alt şeritler sadece rakamla okunur.
    Yatay (kitap açık taranmış) sayfada iki kitap sayfası var: her yarı için ayrı aranır."""
    import pytesseract
    W, H = img.size
    yarilar = [(0, W // 2), (W // 2, W)] if W > H * 1.1 else [(0, W)]
    out = []
    for xa, xb in yarilar:
        kenar = [r for r in rows if (r["top"] / olcek < H * 0.12 or r["bot"] / olcek > H * 0.88)
                 and TS.PAGE_NUM.match(r["text"].strip()) and xa <= (r["x0"] + r["x1"]) / 2 / olcek < xb]
        if kenar:
            continue
        bulunan = []
        for y0, y1 in ((0, int(H * 0.10)), (int(H * 0.90), H)):
            serit = img.crop((xa, y0, xb, y1))
            try:
                d = pytesseract.image_to_data(serit, lang=dil, config="--psm 6", output_type=pytesseract.Output.DICT)
            except Exception:
                continue
            for k, t in enumerate(d["text"]):
                t = (t or "").strip()
                if re.fullmatch(r"\d{1,4}", t) and float(d["conf"][k]) > 60:
                    top = (y0 + d["top"][k]) * olcek
                    bulunan.append({"text": t, "h": d["height"][k] * olcek, "top": top,
                                    "bot": top + d["height"][k] * olcek, "x0": (xa + d["left"][k]) * olcek,
                                    "x1": (xa + d["left"][k] + d["width"][k]) * olcek, "n": 1,
                                    "kalin": False, "blok": -1, "ocr": True})
        out += bulunan[:1]
    return out


def pdf_sayfalari(yol, ilerleme=None):
    """[(satırlar, genişlik, yükseklik, ocr_mu)] ve bilgi sözlüğü."""
    import fitz
    doc = fitz.open(yol)
    n = len(doc)
    sayfalar, ocr = [None] * n, []
    for i, page in enumerate(doc):
        try:
            rows = _katman_satirlari(page)
        except Exception:
            rows = []
        if sum(_harf(r["text"]) for r in rows) < 40 or (_ocr_motoru() == "surya" and _katman_kotu(rows)):
            ocr.append(i)  # metin katmanı yok ya da kötü (eski OCR): sayfa yeniden okunur
        else:
            sayfalar[i] = (rows, page.rect.width, page.rect.height, False)
    bilgi = {"sayfa": n, "ocr": 0, "bozuk_katman": False}
    katman = "\n".join(r["text"] for s in sayfalar if s for r in s[0])
    if TS.katman_bozuk_mu(katman):
        bilgi["bozuk_katman"] = True
        sayfalar, ocr = [None] * n, list(range(n))
    if ocr:
        bilgi["ocr"] = len(ocr)
        dil = os.environ.get("OCR_LANG", "tur")
        kalan_ocr = ocr
        if _ocr_motoru() == "surya":
            okunan = _surya_sayfalar(yol, ocr, ilerleme)
            for i, (rows, w, h) in okunan.items():
                sayfalar[i] = (rows, w, h, True)
            kalan_ocr = [i for i in ocr if i not in okunan]  # Surya yoksa ya da hata verdiyse: Tesseract
            bilgi["ocr_motoru"] = "surya" if okunan else "tesseract"
        if kalan_ocr:
            with ProcessPoolExecutor(max_workers=max(1, os.cpu_count() or 1)) as ex:
                for k, (i, rows, w, h) in enumerate(ex.map(_ocr_sayfa, [(yol, i, dil) for i in kalan_ocr], chunksize=1)):
                    sayfalar[i] = (rows, w, h, True)
                    if ilerleme:
                        ilerleme(f"OCR: sayfa {k + 1}/{len(kalan_ocr)}")
    sayfalar, kaynak = _cift_sayfalari_bol(sayfalar)
    bilgi["kaynak_sayfa"] = kaynak  # her sayfa için (PDF'teki sırası, yarısı: None tek sayfa / 0 ilk / 1 ikinci)
    bilgi["cift_sayfa"] = sum(1 for _, y in kaynak if y == 0)
    ilk_yeni = {}
    for j, (i, _) in enumerate(kaynak):
        ilk_yeni.setdefault(i, j)
    bilgi["kapak_resmi"] = _pdf_kapak(doc)
    meta = doc.metadata or {}
    try:  # PDF yer imleri (bookmarks): [[seviye, başlık, sayfa(1'den)], ...]; çift sayfada yeni sıraya çevrilir
        bilgi["yer_imleri"] = [(int(a), TS.norm(b), ilk_yeni[int(c) - 1] + 1) for a, b, c in doc.get_toc(simple=True)
                               if (b or "").strip() and 1 <= int(c) <= n]
    except Exception:
        bilgi["yer_imleri"] = []
    bilgi["baslik"] = (meta.get("title") or "").strip()
    bilgi["yazar"] = (meta.get("author") or "").strip()
    return sayfalar, bilgi


def ocr_gerekir(yol):
    """0.5.17: dosya ağır (OCR) şeridine mi gider: PDF'te yeniden okunacak sayfalardan (metin katmanı olmayan ya da kötü)
    önbellekte olmayanı varsa. EPUB/DOCX/TXT, iyi katmanlı PDF ve sayfaları önbellekte olan PDF hafif şeritte işlenir.
    Karar pdf_sayfalari'nınkiyle aynı; emin olunamazsa ağır şerit (eski davranış)."""
    if not str(yol).lower().endswith(".pdf"):
        return False
    try:
        import fitz
        doc = fitz.open(yol)
        ocr, katman = [], []
        for i, page in enumerate(doc):
            try:
                rows = _katman_satirlari(page)
            except Exception:
                rows = []
            if sum(_harf(r["text"]) for r in rows) < 40 or (_ocr_motoru() == "surya" and _katman_kotu(rows)):
                ocr.append(i)
            else:
                katman += [r["text"] for r in rows]
        if TS.katman_bozuk_mu("\n".join(katman)):
            ocr = list(range(len(doc)))
        if not ocr:
            return False
        if _ocr_motoru() != "surya":
            return True
        klasor, anahtar = _ocr_onbellek_klasoru(yol), _surya_anahtari()
        return any(_onbellek_oku(klasor, i, anahtar) is None for i in ocr)
    except Exception:
        return True


def _ocr_motoru():
    """OCR_MOTORU: "surya" (varsayılan; Türkçe ve Arapçayı aynı satırda okur) ya da "tesseract"."""
    return os.environ.get("OCR_MOTORU", "surya").strip().lower()


def _katman_kotu(rows):
    """Metin katmanı eski ve kötü bir OCR'dan mı ("tarafmdan", "hir", "varhk"): küçük harfle başlayan kelimelerin
    %30'undan fazlası Zemberek'e göre geçersiz (en az 30 kelime)."""
    kel = re.findall(r"(?<![\w'’])[a-zçğıöşüâîû][a-zçğıöşüâîû]{2,}(?![\w])", " ".join(r["text"] for r in rows))
    if len(kel) < 30:
        return False
    try:
        gecersiz = sum(1 for k in kel if not DZ.gecerli_mi(k))
    except Exception:
        return False
    return gecersiz > 0.3 * len(kel)


_surya = None


def _surya_hazir():
    """Surya modelleri (bir kez yüklenir). Surya kurulu değilse ya da yüklenemezse None."""
    global _surya
    if _surya is None:
        try:
            if os.path.isdir("/data"):
                os.environ.setdefault("MODEL_CACHE_DIR", "/data/modeller/surya")  # modeller kalıcı klasörde
            os.environ.setdefault("TORCH_DEVICE", "cpu")
            import torch
            torch.set_num_threads(max(1, os.cpu_count() or 1))
            from surya.detection import DetectionPredictor
            from surya.recognition import RecognitionPredictor
            _surya = (RecognitionPredictor(), DetectionPredictor())
        except Exception as e:
            print(f"Surya yüklenemedi, Tesseract kullanılacak: {type(e).__name__}: {e}")
            _surya = False
    return _surya or None


_SURYA_DPI = 200


def _surya_satirlari(text_lines, olcek, img=None):
    """Surya satırları -> satır yapısı (PDF birimiyle). <sup>(1)</sup> dipnot işareti olur, <b> kalın bilgisi
    başlık tanımaya gider, öteki biçim işaretleri atılır."""
    rows = []
    for satir in text_lines:
        ham = satir.text or ""
        if img is not None and _AR.search(ham) and _LATIN.search(ham):  # karışık satır: Tesseract konumlarıyla diz
            x0, y0, x1, y1 = satir.bbox
            ham = _karisik_satiri_diz(ham, img.crop((max(0, int(x0) - 4), max(0, int(y0) - 4), int(x1) + 4, int(y1) + 4)))
        t = re.sub(r"<sup>\s*[(\[]?\s*(\d{1,3})\s*[)\]]?\s*</sup>", "\ue000\\1\ue001", ham)
        t = html.unescape(re.sub(r"</?[a-zA-Z][^>]*>", "", t))
        t = _tekrari_at(_surya_harfleri(TS.norm(t)).strip())
        t = re.sub(r"\s+[(\[]\s*$", "", t)  # 0.5.15: sondaki tek açık parantez kalem işareti
        t = _kalem_isaretleri(t)  # 0.5.16
        if len(t.strip()) <= 2 and not _LATIN.search(t) and re.fullmatch(r"[\s\u0660-\u0669\u06f0-\u06f9.\u06d4]+", t):
            continue  # "۰": Arapça rakam/nokta kırıntısı
        if not t:
            continue
        x0, y0, x1, y1 = satir.bbox
        kalin = sum(len(m) for m in re.findall(r"<b>(.*?)</b>", ham)) >= 0.8 * len(re.sub(r"<[^>]+>", "", ham))
        rows.append({"text": t, "h": (y1 - y0) * olcek, "top": y0 * olcek, "bot": y1 * olcek, "x0": x0 * olcek,
                     "x1": x1 * olcek, "n": len(t.split()), "kalin": kalin, "blok": 0, "ocr": True})
    cjk = [r for r in rows if _CJK.search(r["text"])]
    if cjk:
        # Surya boş/lekeli bölgede metin uydurabiliyor ("不可 province.", "Contract"): Çince/Japonca harfli satırlar ve
        # uydurma bölgesindeki (ilk Çince satırın biraz üstünden aşağısı) Arapçasız, kelimelerinin yarısı bile geçerli
        # olmayan satırlar atılır; bölgenin üstündeki gerçek satırlar ("1. KITAP") kalır (0.5.13)
        sinir = min(r["top"] for r in cjk) - 2 * max(r["h"] for r in cjk)
        rows = [r for r in rows if not _CJK.search(r["text"]) and
                (r["bot"] < sinir or _AR.search(r["text"]) or (_gecerli_oran(r["text"]) >= 0.5 and _harf(r["text"]) >= 5))]
        # 0.5.14: uydurma bölgesinde 5 harften kısa satır da atılır ("Carl" Zemberek'e göre geçerli bir kelime)
    elif rows and len(rows) <= 3 and sum(_harf(r["text"]) for r in rows) < 40 and \
            not any(_AR.search(r["text"]) or re.fullmatch(r"\W*\d{1,4}\W*", r["text"]) for r in rows) and (
            not any(_gecerli_oran(r["text"]) >= 0.5 for r in rows) or
            (sum(_harf(r["text"]) for r in rows) < 15 and not any(r["text"].isupper() for r in rows))):
        # boş sayfada uydurma ("Carl"): Arapçası, sayfa numarası olmayan birkaç kısa satır; ya geçerli kelimesi yok ya da
        # tamamı 15 harften az ve büyük harfli başlık değil
        rows = []
    rows = _surya_kucuk_parcalar(rows)
    rows = _surya_ayni_satir(rows)
    for r in rows:
        r["surya"] = True  # 0.5.15
    rows.sort(key=lambda r: (r["top"], r["x0"]))
    return _satir_parcalarini_birlestir(rows)


_AR = re.compile(r"[\u0600-\u06FF\u0750-\u077F\uFB50-\uFDFF\uFE70-\uFEFF]")
_LATIN = re.compile(r"[A-Za-zÇĞİÖŞÜçğıöşüÂÎÛâîû]")
_AYNA = {"(": ")", ")": "(", "[": "]", "]": "[", "«": "»", "»": "«", "﴿": "﴾", "﴾": "﴿"}

# 0.5.11: Surya harf düzeltmeleri. Surya Türkçe "ş" yerine Romence "ș", şapkalı "â" yerine "å"/"ä" yazabiliyor;
# aralıklı başlığın tek harfini Kiril okuyabiliyor ("т AKDIM").
_SURYA_HARF = str.maketrans({"ș": "ş", "Ș": "Ş", "å": "â", "Å": "Â", "ä": "â", "Ä": "Â"})
_KIRIL = re.compile(r"[\u0400-\u04FF]")
_KIRIL_LATIN = str.maketrans({"А": "A", "В": "B", "Е": "E", "К": "K", "М": "M", "Н": "H", "О": "O", "Р": "P",
                              "С": "C", "Т": "T", "Х": "X", "І": "I", "а": "a", "е": "e", "о": "o", "р": "p",
                              "с": "c", "у": "y", "х": "x", "т": "t", "і": "i"})
_CJK = re.compile(r"[\u3040-\u30FF\u3400-\u9FFF\uAC00-\uD7AF]")


def _surya_harfleri(t):
    """Surya'nın Türkçe metinde yaptığı harf karışıklıkları. Kiril harfi yalnız Latin ağırlıklı (ya da tek harflik)
    parçada Latin eşine çevrilir."""
    t = t.translate(_SURYA_HARF)
    if _KIRIL.search(t):
        if len(_LATIN.findall(t)) >= len(_KIRIL.findall(t)) or len(t.strip()) <= 2:
            t = t.translate(_KIRIL_LATIN)
            if len(t.strip()) == 1:
                t = t.upper()
    return t


def _gecerli_oran(t):
    """Satırdaki (3+ harfli) kelimelerin Zemberek'e göre geçerli oranı; kelime yoksa 0."""
    kel = re.findall(r"[^\W\d_]{3,}", t)
    if not kel:
        return 0.0
    try:
        return sum(1 for k in kel if DZ.gecerli_mi(k) or DZ._kelime_mi(k)) / len(kel)
    except Exception:
        return 1.0


def _surya_kucuk_parcalar(rows):
    """Surya aynı satırdaki kısa parçayı (aralıklı başlığın ilk harfi "T", madde harfi "A —") ayrı ve daha alçak bir
    satır olarak verebiliyor; yükseklik oranı şartı yüzünden _satir_parcalarini_birlestir onu tutmaz ve parça tek
    başına kalıp silinir. Kısa parça (en çok 4 işaret) dikey ortası komşu satırın içindeyse ve aradaki yatay boşluk bir
    satır yüksekliğinden azsa soldan sağa birleşir. Tek büyük harf, büyük harfli kelimeye bitişik yazılır (TAKDIM)."""
    out = [dict(r) for r in rows]
    degisti = True
    while degisti:
        degisti = False
        for a in out:
            ta = a["text"].strip()
            if not ta or len(ta.replace(" ", "")) > 4:
                continue
            orta = (a["top"] + a["bot"]) / 2
            for b in out:
                tb = b["text"].strip()
                if b is a or len(tb.replace(" ", "")) <= len(ta.replace(" ", "")) or not b["top"] <= orta <= b["bot"]:
                    continue
                hb = b["bot"] - b["top"]
                if a["x1"] <= b["x0"] + 1 and b["x0"] - a["x1"] <= hb:
                    ilk = tb.split()[0]
                    bitisik = len(ta) == 1 and ta.isupper() and ilk.isalpha() and ilk.isupper()
                    b["text"] = ta + ("" if bitisik else " ") + tb
                    b["x0"] = min(a["x0"], b["x0"])
                elif b["x1"] <= a["x0"] + 1 and a["x0"] - b["x1"] <= hb:
                    b["text"] = tb + " " + ta
                    b["x1"] = max(a["x1"], b["x1"])
                else:
                    continue
                b["n"] = len(b["text"].split())
                b["kalin"] = b["kalin"] and a["kalin"]
                out.remove(a)
                degisti = True
                break
            if degisti:
                break
    return out


# 0.5.13: Surya'ya özel kurallar
def _tekrari_at(t):
    """Surya bazen satırın bir bölümünü iki kez yazıyor ("senelerinde İstanbul'da basılan ( … senelerinde İstanbul'da
    basılan ("): 3 ya da daha çok kelimelik dizinin ikinci geçişi atılır."""
    k = t.split()
    for n in range(min(8, len(k) // 2), 2, -1):
        for i in range(len(k) - 2 * n + 1):
            for j in range(i + n, len(k) - n + 1):
                if k[i:i + n] == k[j:j + n] and _LATIN.search(" ".join(k[i:i + n])):
                    return _tekrari_at(" ".join(k[:j] + k[j + n:]))
    return t


def _surya_ayni_satir(rows):
    """Surya aynı satırdaki Arapça ibareyi ve Türkçeyi ayrı satır verebiliyor; Arapça satır yüksek olduğu için
    _satir_parcalarini_birlestir onları birleştirmez ve yukarıdan aşağı sıralamada Arapça öne geçer ("( … )" +
    "olsun) : Allah'ın Resûlü"). Biri Arapçalıysa, dikey ortası ötekinin içinde ve aradaki boşluk bir satır
    yüksekliğinden azsa soldan sağa birleşir (iki sütunlu sayfada sütunlar arası daha geniştir)."""
    out = [dict(r) for r in rows]
    degisti = True
    while degisti:
        degisti = False
        for a in out:
            for b in out:
                if b is a or not (_AR.search(a["text"]) or _AR.search(b["text"])):
                    continue
                if _AR.search(a["text"]) and _AR.search(b["text"]) and not (_LATIN.search(a["text"]) or _LATIN.search(b["text"])):
                    continue
                ya, yb = (a["top"] + a["bot"]) / 2, (b["top"] + b["bot"]) / 2
                if not (b["top"] <= ya <= b["bot"] or a["top"] <= yb <= a["bot"]):
                    continue
                sol, sag = (a, b) if a["x0"] <= b["x0"] else (b, a)
                if not (-2 <= sag["x0"] - sol["x1"] <= max(a["h"], b["h"])):
                    continue
                sol["text"] = sol["text"].rstrip() + " " + sag["text"].lstrip()
                sol["x1"] = max(sol["x1"], sag["x1"])
                sol["top"], sol["bot"] = min(a["top"], b["top"]), max(a["bot"], b["bot"])
                sol["h"] = max(a["h"], b["h"])
                sol["n"] = len(sol["text"].split())
                sol["kalin"] = a["kalin"] and b["kalin"]
                out.remove(sag)
                degisti = True
                break
            if degisti:
                break
    return out


# 0.5.15: Türkçe harfi düşmüş büyük harfli kelime ("BESINCI" -> "BEŞİNCİ"): tek geçerli aday varsa o
_TR_ESLER = {"I": "İ", "S": "Ş", "C": "Ç", "G": "Ğ", "O": "Ö", "U": "Ü", "i": "ı", "s": "ş", "c": "ç", "g": "ğ",
             "o": "ö", "u": "ü"}


def _gecerli_kelime(k):
    try:
        return bool(DZ.gecerli_mi(k) or DZ._kelime_mi(k))
    except Exception:
        return True


def _turkce_kelime_onar(k):
    """Geçersiz kelimede düşmüş Türkçe harfleri dener (en çok 6 yer); geçerli adaylardan en çok harfi düzelten tek
    aday varsa onu döndürür ("BESINCI" -> "BEŞİNCİ")."""
    if len(k) < 3 or _gecerli_kelime(k):
        return k
    yerler = [i for i, c in enumerate(k) if c in _TR_ESLER][:6]
    if not yerler:
        return k
    import itertools
    adaylar = set()
    for secim in itertools.product((False, True), repeat=len(yerler)):
        if not any(secim):
            continue
        h = list(k)
        for i, deg in zip(yerler, secim):
            if deg:
                h[i] = _TR_ESLER[h[i]]
        a = "".join(h)
        if _gecerli_kelime(a):
            adaylar.add((sum(secim), a))
    if not adaylar:
        return k
    en = max(n for n, _ in adaylar)
    enler = [a for n, a in adaylar if n == en]  # Zemberek gevşek: "BESİNCİ" de geçer; en çok harfi düzelten aday
    return enler[0] if len(enler) == 1 else k


def _buyuk_baslik_onar(t):
    """Büyük harfli başlıkta harfi düşmüş kelimeler onarılır (uzunluk aynı kalır: sayfa konumları geçerli)."""
    return re.sub(r"[^\W\d_]{3,}", lambda m: _turkce_kelime_onar(m.group()) if m.group().isupper() else m.group(), t)


# 0.5.16: kalemle çizilmiş köşeli parantezler ve çizgiler Surya'da "∫", "|", "[", "]", "/" olarak okunur
def _kalem_isaretleri(t):
    t = t.replace("∫", " ")
    if "[" not in t:
        t = re.sub(r"(?<=[.,:;!?\"”»)])\s*\]+", "", t)  # "yaptırmak.]", "tamamlar. ]"
        t = re.sub(r"\s+\]+(?=\s|$)", "", t)
    if "]" not in t:
        t = re.sub(r"^\s*\[+\s*(?=[A-ZÇĞİÖŞÜÂÎÛ0-9\"“])", "", t)
    t = re.sub(r"^\s*\|+\s*(?=[A-ZÇĞİÖŞÜÂÎÛ0-9\"“])", "", t)  # "|Üçüncü kısım"
    t = re.sub(r"\s+\|+\s*$", "", t)
    t = re.sub(r"(?<=[^\W\d_])/(?=[.,;:])", "", t)  # "oruç/."
    return re.sub(r"\s{2,}", " ", t).strip()


def _harfli_madde_basligi(t):
    """"A — ZAMAN", "C - RIZIK MESELESİ", "B — İRTİKÂ (Yücelme) TABİRİ": harfle numaralanmış, büyük harfli ara başlık."""
    t = SAYFA_ISARET.sub("", t).strip()
    m = re.match(r"^[A-ZÇĞİÖŞÜ]\s?[—–-]\s?(\S.*)$", t)
    if not m or len(t) > 70 or t.endswith((".", ",", ";")):
        return False
    govde = re.sub(r"\([^()]*\)", " ", m.group(1))
    harf = [c for c in govde if c.isalpha()]
    return len(harf) >= 3 and all(c.isupper() for c in harf)


_AR_PARCA = re.compile(r"[\u0600-\u06FF\u0750-\u077F\uFB50-\uFDFF\uFE70-\uFEFF\u064B-\u065F﴿﴾]"
                       r"(?:[\s\u0600-\u06FF\u0750-\u077F\uFB50-\uFDFF\uFE70-\uFEFF\u064B-\u065F﴿﴾:.]*"
                       r"[\u0600-\u06FF\u0750-\u077F\uFB50-\uFDFF\uFE70-\uFEFF\u064B-\u065F﴿﴾])?")


def _karisik_satiri_diz(metin, kesit=None):
    """Surya, Arapça ile Türkçenin aynı satırda olduğu satırı sağdan sola dizebiliyor ("( فليرتقوا … ) Kur'an-ı
    Kerim'in bir emri olan (") ve kelime kutuları da güvenilmez. Metin Surya'dan, yer Tesseract'tan: satırın görüntüsü
    (kesit) Türkçe okunur, her Türkçe parça ilk kelimeleriyle Tesseract kelimesine eşlenip yeri bulunur; Arapça parça
    Tesseract'ın eşleşmeyen (okuyamadığı) bölgesine konur. Parçalar soldan sağa dizilir, Arapça parça parantez içine
    alınır, kenarda ters kalmış parantez/noktalama ve satır içi tekrar ("(Fussilet 41/10)" iki kez) atılır. Bir parçanın
    yeri bulunamazsa satıra dokunulmaz."""
    if kesit is None or not (_AR.search(metin) and _LATIN.search(metin)):
        return metin
    import difflib
    import pytesseract
    parcalar, son = [], 0  # (tür, yazı)
    for m in _AR_PARCA.finditer(metin):
        if not _AR.search(m.group()):
            continue
        if metin[son:m.start()].strip():
            parcalar.append(("L", metin[son:m.start()]))
        parcalar.append(("A", m.group()))
        son = m.end()
    if metin[son:].strip():
        parcalar.append(("L", metin[son:]))
    if not any(t == "A" for t, _ in parcalar):
        return metin
    try:
        d = pytesseract.image_to_data(kesit, lang="tur", config="--psm 7", output_type=pytesseract.Output.DICT)
    except Exception:
        return metin
    tk = [(d["left"][k], (d["text"][k] or "").strip()) for k in range(len(d["text"])) if (d["text"][k] or "").strip()]
    gen = [d["width"][k] for k in range(len(d["text"])) if (d["text"][k] or "").strip()]
    # 0.5.14: şapka ve nokta farkı yok sayılır ("Teala" = "Teâlâ", "ADİ" = "adi"); her kelime en iyi adayına eşlenir
    sade = lambda w: "".join(c for c in unicodedata.normalize("NFKD", w.lower().replace("ı", "i"))
                             if c.isalnum() and not unicodedata.combining(c))
    eslesen, yer = set(), []
    parcalar = [(t, y) for t, y in parcalar if t == "A" or re.search(r"[^\W\d_]{2}", y)]  # yalnız işaret: at
    for tur, yazi in parcalar:
        if tur == "A":
            yer.append(None)
            continue
        adaylar = []
        for w in [w for w in yazi.split() if len(sade(w)) >= 2]:
            en_iyi = None
            for k, (x, t) in enumerate(tk):
                if k in eslesen or k in {a for _, a in adaylar} or len(sade(t)) < 2:
                    continue
                oran = difflib.SequenceMatcher(None, sade(w), sade(t)).ratio()
                if oran >= 0.75 and (en_iyi is None or oran > en_iyi[0]):
                    en_iyi = (oran, x, k)
            if en_iyi:
                adaylar.append((en_iyi[1], en_iyi[2]))
        if not adaylar:
            return metin  # Türkçe parçanın yeri bulunamadı: dokunma
        eslesen.update(k for _, k in adaylar)
        yer.append(min(x for x, _ in adaylar))
    bos = [x for k, (x, t) in enumerate(tk) if k not in eslesen]
    if sum(1 for t, _ in parcalar if t == "A") == 1 and eslesen:
        # 0.5.13: tek Arapça parça. Tesseract'ın okuyamadığı kelimeler Türkçenin arasında da kalabilir (ortancayı
        # yanıltır): önce Türkçe kelimelerin kapladığı yerin DIŞINDAKİ eşleşmeyen kelimeler; hiç yoksa satırın
        # mürekkepli kısmında Türkçenin boş bıraktığı en geniş bölge
        dolu = sorted((tk[k][0], tk[k][0] + gen[k]) for k in eslesen)
        disarda = [x for k, (x, t) in enumerate(tk) if k not in eslesen and not any(a - 2 <= x <= b for a, b in dolu)]
        if disarda:
            bos = disarda
        else:
            try:
                kutu = kesit.convert("L").point(lambda v: 255 if v < 128 else 0).getbbox()
            except Exception:
                kutu = None
            sol_kenar, sag_kenar = (kutu[0], kutu[2]) if kutu else (0, getattr(kesit, "width", 0))
            aralik, bas = [], sol_kenar
            for a, b in dolu:
                if a > bas:
                    aralik.append((a - bas, (bas + a) / 2))
                bas = max(bas, b)
            if sag_kenar > bas:
                aralik.append((sag_kenar - bas, (bas + sag_kenar) / 2))
            if aralik:
                bos = [max(aralik)[1]]
    for n, (tur, _) in enumerate(parcalar):
        if tur == "A":
            if not bos:
                return metin
            yer[n] = st.median(bos)
    sirali = [p for _, p in sorted(zip(yer, parcalar), key=lambda z: z[0])]
    out = []
    for tur, yazi in sirali:
        yazi = yazi.strip()
        if tur == "A":
            ic = re.sub(r"[()﴿﴾]", " ", yazi).strip(" :.،")  # kenardaki Arapça virgül ve içteki fazla parantez
            ic = re.sub(r"\s+", " ", ic)
            out.append(("﴿ " + ic + " ﴾") if "﴿" in yazi or "﴾" in yazi else ("( " + ic + " )"))
        else:
            yazi = yazi.replace("،", ",")
            for tekrar in set(re.findall(r"\([^()]{2,40}\)", yazi)):  # satır içi tekrar: sonuncusu kalır
                while yazi.count(tekrar) > 1:
                    yazi = yazi.replace(tekrar, "", 1).strip()
            # sağdan sola dizilmenin kenara attığı işaretler: baştaki nokta cümle sonudur (sona), öteki baştaki işaret
            # ve sondaki "(" atılır; eşi olmayan "»" sondan başa "«" olur
            bas = re.match(r"^[\s:.,;)]*", yazi).group()
            yazi = yazi[len(bas):].rstrip(" (")
            if re.match(r"^-[^\W\d_]", yazi) and not yazi.endswith("-"):  # ters dizilmiş satır sonu tiresi: "-diye … tanın"
                yazi = yazi[1:] + "-"
            if "." in bas and not yazi.endswith((".", "»", ")")):
                yazi += "."
            if yazi.endswith("»") and "«" not in yazi:
                yazi = "«" + yazi[:-1].rstrip()
            if yazi.count("(") > yazi.count(")") and yazi.startswith("("):
                yazi = yazi[1:].lstrip()
            elif yazi.count(")") > yazi.count("("):  # ".melekleri)dir (" -> "(melekleri)dir"
                yazi = "(" + yazi
            if yazi:
                out.append(yazi)
    return re.sub(r"\s+", " ", " ".join(out)).strip()


def _surya_anahtari():
    """Önbellek anahtarı: Surya sürümü ve çözünürlük değişirse eski kayıt kullanılmaz."""
    try:
        from importlib.metadata import version
        surum = version("surya-ocr")
    except Exception:
        surum = "?"
    return f"surya-{surum}-dpi{_SURYA_DPI}"


def _ocr_onbellek_klasoru(yol):
    """0.5.12: Surya'nın ham satırları sayfa sayfa saklanır; kod düzeltilip kitap yeniden işlenince OCR tekrarlanmaz.
    Klasör dosyanın İÇERİĞİNİN özetiyle adlandırılır (aynı dosya yeniden yüklense de bulunur). OCR_ONBELLEK=0 kapatır."""
    kok = os.environ.get("OCR_ONBELLEK", os.path.join(os.environ.get("DATA_DIR", "/data"), "ocr_onbellek"))
    if kok.strip().lower() in ("", "0", "kapali", "kapalı"):
        return None
    try:
        h = hashlib.sha1()
        with open(yol, "rb") as f:
            for parca in iter(lambda: f.read(1 << 20), b""):
                h.update(parca)
        return os.path.join(kok, h.hexdigest()[:20])
    except OSError:
        return None


def _onbellek_oku(klasor, i, anahtar):
    """Sayfanın kayıtlı ham satırları [(metin, (x0, y0, x1, y1))] ya da None (yok, bozuk ya da başka sürüm)."""
    if not klasor:
        return None
    try:
        import json
        with open(os.path.join(klasor, f"s{i + 1:04d}.json"), encoding="utf-8") as f:
            kayit = json.load(f)
        if kayit.get("anahtar") != anahtar:
            return None
        return [(str(t), tuple(float(v) for v in b)) for t, b in kayit["satirlar"]]
    except Exception:
        return None


def _onbellek_yaz(klasor, i, anahtar, text_lines, pdf_adi=""):
    """Ham Surya satırları yazılır (geçici dosya + yeniden adlandırma: yarım dosya kalmaz). Yazılamazsa (salt okunur
    klasör, dolu disk) sessizce geçilir: önbellek yalnız hız içindir."""
    if not klasor:
        return
    try:
        import json
        os.makedirs(klasor, exist_ok=True)
        if pdf_adi and not os.path.exists(os.path.join(klasor, "kaynak.txt")):
            with open(os.path.join(klasor, "kaynak.txt"), "w", encoding="utf-8") as f:
                f.write(pdf_adi + "\n")
        kayit = {"anahtar": anahtar, "sayfa": i + 1,
                 "satirlar": [[s.text or "", [round(float(v), 2) for v in s.bbox]] for s in text_lines]}
        hedef = os.path.join(klasor, f"s{i + 1:04d}.json")
        with open(hedef + ".tmp", "w", encoding="utf-8") as f:
            json.dump(kayit, f, ensure_ascii=False)
        os.replace(hedef + ".tmp", hedef)
    except Exception as e:
        print(f"OCR önbelleğine yazılamadı (sayfa {i + 1}): {type(e).__name__}: {e}")


def _surya_sayfalar(yol, sayfalar, ilerleme=None, grup=4):
    """Sayfaları Surya ile okur (tek süreçte, dörder sayfa). {sayfa: (satırlar, genişlik, yükseklik)}; okunamayan
    sayfa sonuçta yer almaz (Tesseract'a kalır). 0.5.12: önbellekteki sayfalar Surya'ya gitmez; ham satırlar
    önbellekten okunup aynı işlemden (_surya_satirlari) geçer, yani kod düzeltmeleri bunlara da uygulanır."""
    if not sayfalar:
        return {}
    import types
    import fitz
    from PIL import Image
    doc = fitz.open(yol)
    olcek = 72 / _SURYA_DPI
    klasor = _ocr_onbellek_klasoru(yol)
    anahtar = _surya_anahtari()

    def resim(i):
        pix = doc[i].get_pixmap(dpi=_SURYA_DPI)
        return Image.frombytes("RGB", (pix.width, pix.height), pix.samples)

    out, kalan = {}, []
    for i in sayfalar:
        ham = _onbellek_oku(klasor, i, anahtar)
        if ham is None:
            kalan.append(i)
            continue
        satirlar = [types.SimpleNamespace(text=t, bbox=b) for t, b in ham]
        try:  # karışık (Arapça + Türkçe) satır varsa dizme için sayfa resmi gerekir
            img = resim(i) if any(_AR.search(s.text) and _LATIN.search(s.text) for s in satirlar) else None
            out[i] = (_surya_satirlari(satirlar, olcek, img), doc[i].rect.width, doc[i].rect.height)
        except Exception as e:
            print(f"Önbellekteki sayfa {i + 1} işlenemedi, yeniden okunacak: {type(e).__name__}: {e}")
            kalan.append(i)
    onbellekten = len(out)
    if ilerleme and onbellekten:
        ilerleme(f"OCR önbellekten: {onbellekten}/{len(sayfalar)} sayfa")
    if not kalan:
        return out
    model = _surya_hazir()
    if not model:
        return out
    rec, det = model
    ek = f" (önbellekten {onbellekten})" if onbellekten else ""
    for k in range(0, len(kalan), grup):
        parca = kalan[k:k + grup]
        try:
            resimler = [resim(i) for i in parca]
            sonuc = rec(resimler, det_predictor=det, sort_lines=True, math_mode=False)
            for i, s, img in zip(parca, sonuc, resimler):
                _onbellek_yaz(klasor, i, anahtar, s.text_lines, os.path.basename(yol))
                out[i] = (_surya_satirlari(s.text_lines, olcek, img), doc[i].rect.width, doc[i].rect.height)
        except Exception as e:
            print(f"Surya {parca} sayfalarında hata, Tesseract'a kalıyor: {type(e).__name__}: {e}")
        if ilerleme:
            ilerleme(f"OCR (Surya): sayfa {min(k + grup, len(kalan))}/{len(kalan)}{ek}")
    return out


def _cilt_arasi(rows, w, h):
    """Yatay sayfada iki kitap sayfası yan yana mı (kitap açık taranmış): hiçbir satırın üstünden geçmediği, ortaya
    yakın boş dikey şeridin ortası; yoksa None."""
    if w <= h * 1.1:
        return None
    yazili = [r for r in rows if _harf(r["text"]) >= 3]
    if len(yazili) < 6:
        return None
    bos = [g for g in range(int(w * 0.3), int(w * 0.7) + 1)
           if not any(r["x0"] < g - 1 and r["x1"] > g + 1 for r in yazili)]
    seritler, bas = [], None
    for k, g in enumerate(bos):
        if bas is None:
            bas = g
        if k + 1 == len(bos) or bos[k + 1] != g + 1:
            seritler.append((bas, g))
            bas = None
    seritler = [(a, b) for a, b in seritler if b - a >= 6]
    if not seritler:
        return None
    a, b = min(seritler, key=lambda s: abs((s[0] + s[1]) / 2 - w / 2) - (s[1] - s[0]) / 4)
    g = (a + b) / 2
    sol, sag = sum(1 for r in yazili if r["x1"] <= g), sum(1 for r in yazili if r["x0"] >= g)
    return g if sol >= 3 and sag >= 3 else None


def _cift_sayfalari_bol(sayfalar):
    """Kitap açık taranmış (her PDF sayfasında iki kitap sayfası) PDF'lerde her yatay sayfa cilt arasından ikiye
    bölünür; sağ yarının satırları sola kaydırılır. Sıra: Latin yazılı kitapta sol-sağ, Arapça yazılıda sağ-sol.
    Döndürür: (yeni sayfa listesi, [(PDF'teki sırası, yarısı)])."""
    yeni, kaynak = [], []
    for i, (rows, w, h, ocr) in enumerate(sayfalar):
        g = _cilt_arasi(rows, w, h)
        if g is None:
            yeni.append((rows, w, h, ocr))
            kaynak.append((i, None))
            continue
        sira = lambda r: (r["top"], r["x0"])  # tek sütunlu kitap sayfası: yukarıdan aşağı (Tesseract iki sayfanın
        # satırlarını birleştirip bölünmüşse sağ yarının satırları karışık sırada gelir)
        sol = sorted((r for r in rows if (r["x0"] + r["x1"]) / 2 < g), key=sira)
        sag = sorted((dict(r, x0=r["x0"] - g, x1=r["x1"] - g) for r in rows if (r["x0"] + r["x1"]) / 2 >= g), key=sira)
        metin = " ".join(r["text"] for r in rows)
        sagdan = len(re.findall(r"[\u0600-\u06FF]", metin)) > len(re.findall(r"[A-Za-zÇĞİÖŞÜçğıöşü]", metin))
        yarilar = [(sag, w - g), (sol, g)] if sagdan else [(sol, g), (sag, w - g)]
        for k, (rs, ww) in enumerate(yarilar):
            yeni.append((rs, ww, h, ocr))
            kaynak.append((i, k))
    return yeni, kaynak


def _pdf_kapak(doc):
    """Kitabın kendi kapağı: ilk 3 sayfadan ilk dolu (boş/düz renk olmayan) sayfanın görüntüsü (JPEG)."""
    try:
        from PIL import Image, ImageStat
        for i in range(min(3, len(doc))):
            page = doc[i]
            pix = page.get_pixmap(dpi=max(72, int(72 * 1600 / max(page.rect.height, 1))))
            img = Image.open(io.BytesIO(pix.tobytes("png"))).convert("RGB")
            if max(ImageStat.Stat(img.convert("L")).stddev) < 8:
                continue  # boş ya da düz renk sayfa
            b = io.BytesIO()
            img.save(b, "JPEG", quality=85, optimize=True)
            return (b.getvalue(), "image/jpeg")
    except Exception:
        pass
    return None


def _harf_farki(a, b):
    """İki yazı arasındaki harf farkı (Levenshtein)."""
    onceki = list(range(len(b) + 1))
    for i, ca in enumerate(a, 1):
        simdi = [i]
        for j, cb in enumerate(b, 1):
            simdi.append(min(onceki[j] + 1, simdi[j - 1] + 1, onceki[j - 1] + (ca != cb)))
        onceki = simdi
    return onceki[-1]


def _icindekiler_basligi(t):
    """'İÇİNDEKİLER', 'İçindekiler:', 'FİHRİST' (büyük İ'nin küçültülmesindeki birleşik nokta dahil); bozuk OCR
    ('iONDEKİ LER') en çok 2 harf farkla."""
    t = "".join(c for c in unicodedata.normalize("NFKD", t.strip(" .:")) if not unicodedata.combining(c))
    t = t.lower().replace("ı", "i")
    if t in ("icindekiler", "fihrist", "contents", "table of contents"):
        return True
    sik = re.sub(r"[\W\d_]", "", t)
    return 8 <= len(sik) <= 14 and _harf_farki(sik, "icindekiler") <= 2


def _icindekiler_sayfasi(rows):
    """Basılı içindekiler sayfası mı: başlığında İçindekiler/Fihrist ya da satırlarının çoğu sayfa numarasıyla bitiyor."""
    if not rows:
        return False
    if any((TS.ICINDEKILER.search(r["text"]) and len(r["text"]) < 40) or _icindekiler_basligi(r["text"]) for r in rows[:4]):
        return True
    return len(rows) >= 6 and sum(1 for r in rows if TS.NOKTALI.search(r["text"])) >= 0.6 * len(rows)


# ======================= PDF: sayfa düzeyinde ayıklama =======================
_UST_BILGI_NO = re.compile(r"^(\d{1,4})\s+\S.{0,70}$|^.{1,70}\S\s+(\d{1,4})$")


def _sayfa_no_ve_kenar(rows, h):
    """Basılı sayfa numarasını bulur, sayfa numarası satırlarını çıkarır. (no, kalan satırlar)"""
    no, kalan, tahmin = None, [], None
    # 0.5.15: Surya'da üst bilgi her sayfada farklı okunabilir ("Dinde Kırk Prensip", "Dunde Kirk Prensip") ve tekrar eden
    # satır diye tanınmaz: sayfanın üstündeki numara satırıyla aynı hizadaki kısa satırlar üst bilgidir (kalemle
    # eklenmiş not da dahil). Numara satırının kendisi aşağıda sayfa numarası olarak okunur.
    ust_no = [r for r in rows if r.get("surya") and r["top"] < h * 0.12 and r["h"] >= 6 and
              re.fullmatch(r"\W*\d{1,4}\W*", _rakam(r["text"]))]
    if ust_no:
        n = ust_no[0]
        bant = (n["top"] - (n["bot"] - n["top"]), n["bot"] + (n["bot"] - n["top"]))
        rows = [r for r in rows if r is n or not (r.get("surya") and len(r["text"]) <= 80 and
                                                  bant[0] <= (r["top"] + r["bot"]) / 2 <= bant[1])]
    for r in rows:
        kenar = r["top"] < h * 0.12 or r["bot"] > h * 0.88
        t = _rakam(r["text"].strip())
        if kenar and TS.PAGE_NUM.match(t):
            no = no or re.sub(r"[^\divxlcdm]", "", t.lower())
            continue
        if kenar and r["n"] <= 12:
            m = _UST_BILGI_NO.match(t)
            if m and tahmin is None:
                tahmin = m.group(1) or m.group(2)
        kalan.append(r)
    return no or tahmin, kalan  # 0.5.15: kesin numara satırı öncelikli


def _tekrar_edenleri_at(sayfa_satirlari):
    """Sayfaların ilk/son satırlarında tekrar eden üst/alt bilgiler (kitap adı, bölüm adı)."""
    anahtar = lambda t: re.sub(r"[\d\W]+", " ", t).strip().lower()
    sayac = collections.Counter()
    for rows in sayfa_satirlari:
        for r in rows[:2] + rows[-2:]:
            if len(r["text"]) < 90:
                sayac[anahtar(r["text"])] += 1
    esik = max(3, int(len(sayfa_satirlari) * 0.3))
    tekrar = {k for k, c in sayac.items() if c >= esik and k}
    out = []
    for rows in sayfa_satirlari:
        ilk_son = {id(r) for r in rows[:2] + rows[-2:]}
        out.append([r for r in rows if not (id(r) in ilk_son and anahtar(r["text"]) in tekrar)])
    return out


def _eksik_numaralari_doldur(nolar):
    """Numarası okunamayan sayfalar: arap rakamlı sayfalardan kayma (basılı - pdf sırası) bulunup doldurulur."""
    kayma = collections.Counter(int(n) - i for i, n in enumerate(nolar) if n and n.isdigit())
    if not kayma:
        return nolar
    k, adet = kayma.most_common(1)[0]
    if adet < 2:
        return nolar
    out = []
    for i, n in enumerate(nolar):
        if n and n.isdigit():
            # yakın sayfalardaki kaymaya göre denetle (kitap ortasında boş/eksik sayfa kaymayı değiştirebilir)
            komsu = [int(m) - j for j, m in enumerate(nolar) if m and m.isdigit() and 0 < abs(j - i) <= 8]
            yerel = collections.Counter(komsu).most_common(1)[0][0] if komsu else k
            if abs(int(n) - i - yerel) > 3:
                n = None  # yanlış okunmuş (ör. 237 yerine 2877)
        out.append(n if n else (str(i + k) if i + k >= 1 else None))
    return out


# ======================= PDF: paragraflar ve başlıklar =======================
def _govde_boyu(sayfa_satirlari, ocr):
    """Gövde yazısının boyu. Katmanda en sık punto; OCR'da (yükseklik harflere göre oynar) ağırlıklı ortanca."""
    rows = [r for rs in sayfa_satirlari for r in rs if r.get("ocr") == ocr and r["n"] >= 4]
    if not rows:
        return None
    if not ocr:
        say = collections.Counter()
        for r in rows:
            say[round(r["h"] * 2) / 2] += len(r["text"])
        return say.most_common(1)[0][0]
    agir = sorted((r["h"], len(r["text"])) for r in rows)
    yari, top = sum(a for _, a in agir) / 2, 0
    for h, a in agir:
        top += a
        if top >= yari:
            return h
    return agir[-1][0]


def _baslik_mi(r, govde, genislik, kalin_oran):
    t = r["text"].strip()
    if len(t) > 90 or len(t) < 2 or t.endswith((",", ";")) or UST.search(re.sub(r"\ue000\d{1,3}\ue001\s*$", "", t)) or \
            re.search(r"[^\W\d_]-$", t):  # 0.5.15; 0.5.16: tireyle biten (kelimesi bölünmüş) satır başlık değil
        return False
    if TS.ICERIK_BASLIK.match(t) or BOLUM_NO.match(t) or _harfli_madde_basligi(t):
        return True
    arapcali = bool(_AR.search(t))  # Arapça harfler satırı yüksek gösterir: büyük punto sanılmasın
    if r["h"] >= govde * (1.4 if r.get("ocr") else 1.15) and not arapcali:
        return True
    ortada = abs((r["x0"] + r["x1"]) / 2 - genislik / 2) < genislik * 0.08 and (r["x1"] - r["x0"]) < genislik * 0.7
    if r["kalin"] and kalin_oran < 0.3 and not t.endswith(".") and (ortada or len(t) < 60):
        return True
    harfler = [c for c in t if c.isalpha()]
    if bool(harfler) and all(c.isupper() for c in harfler) and len(t) <= 50 and ortada:
        return True
    if r.get("ocr") and len(harfler) >= 8 and all(c.isupper() for c in harfler) and 2 <= len(t.split()) <= 6 and \
            len(t) <= 50 and not t.endswith(END_PUNCT) and not _AR.search(t):
        return True  # 0.5.13: "RÜYÂDA GÖRÜLMESİ" (iki satırlı başlığın ikinci satırı, ortalanmamış)
    # OCR: iki yandan da içeride, ortalanmış, noktayla bitmeyen kısa satır
    return bool(r.get("ocr")) and ortada and r["x0"] > genislik * 0.25 and len(t) <= 70 and not t.endswith(END_PUNCT) \
        and r["h"] >= govde * 1.1 and not arapcali


_KENAR_BAS = re.compile(r"^\((\d{1,3})\)\s+(?=\S)")
_KENAR_SON = re.compile(r"(?<=\S)\s+\((\d{1,3})\)\s*$")
_KENAR_TEK = re.compile(r"^\s*\((\d{1,3})\)\s*$")  # tek başına satır: "(9)"


def _kenar_numarasi_dizisi(sayfa_satirlari):
    """OCR katmanı kenar numarasını (aslın sayfa numarası gibi) satırın başına/sonuna yazmışsa ayıklar: "millet- (17)",
    "(18) edilebilen". Güvenlik: satır sınırındaki numaralar kitap boyunca düzenli artan bir dizi oluşturmalı."""
    bulunan = []
    for rows in sayfa_satirlari:
        for r in rows:
            m = _KENAR_TEK.search(r["text"]) or _KENAR_BAS.search(r["text"]) or _KENAR_SON.search(r["text"])
            if m:
                bulunan.append(int(m.group(1)))
    if len(bulunan) < 8:
        return sayfa_satirlari
    artan = sum(1 for a, b in zip(bulunan, bulunan[1:]) if 0 < b - a <= 3)
    if artan < 0.7 * (len(bulunan) - 1):
        return sayfa_satirlari  # düzenli dizi değil: metnin kendi numaraları olabilir, dokunulmaz
    for rows in sayfa_satirlari:
        for r in rows:
            yeni = "" if _KENAR_TEK.search(r["text"]) else _KENAR_SON.sub("", _KENAR_BAS.sub("", r["text"]))
            if yeni != r["text"]:
                r["text"] = yeni
                r["n"] = len(yeni.split())
    return [[r for r in rows if r["text"].strip()] for rows in sayfa_satirlari]


def _kenar_numaralari_at(rows):
    """Metin sütununun dışında (sağ ya da sol kenarda) duran "(17)", "17" gibi numaralar: kenar notu (ör. aslın sayfa
    numarası); metne karışmasın."""
    govde = [r for r in rows if r["n"] >= 4]
    if len(govde) < 3:
        return rows
    sol = min(r["x0"] for r in govde)
    sag = max(r["x1"] for r in govde)
    return [r for r in rows if not (re.fullmatch(r"\(?\s*\d{1,4}\s*\)?", r["text"].strip()) and
                                    (r["x0"] >= sag - 2 or r["x1"] <= sol + 2))]


def _sayfa_paragraflari(rows, genislik, govde, kalin_oran, bas_ayri=False):
    """Bir sayfanın ana satırları -> [(tür, metin, boy)] (tür 'b' başlık ya da 'p').
    bas_ayri: art arda başlık satırları birleştirilmez (içindekilerden fihrist: hangisinin ne olduğunu o söyler)."""
    rows = _kenar_numaralari_at(rows)
    if not rows:
        return []
    govde_satir = [r for r in rows if abs(r["h"] - govde) < govde * 0.15]
    sol = collections.Counter(round(r["x0"]) for r in govde_satir).most_common(1)[0][0] if govde_satir else min(r["x0"] for r in rows)
    sag = max((r["x1"] for r in govde_satir), default=max(r["x1"] for r in rows))
    araliklar = [b["top"] - a["top"] for a, b in zip(rows, rows[1:]) if 0 < b["top"] - a["top"] < govde * 3]
    aralik = st.median(araliklar) if araliklar else govde * 1.4
    out, cur, onceki = [], [], None
    for r in rows:
        baslik = _baslik_mi(r, govde, genislik, kalin_oran)
        yeni = True
        if onceki is not None and cur:
            if baslik and cur[0][0] == "b" and r["top"] - onceki["top"] < aralik * 2.2 and \
                    abs(r["h"] - onceki["h"]) < (govde * 0.3 if r.get("ocr") else 0.5) and \
                    not (r.get("ocr") and BOLUM_NO.match(onceki["text"].strip())):
                yeni = False  # iki satıra bölünmüş başlık
            elif not baslik and cur[0][0] == "p":
                girinti = r["x0"] > sol + govde * 0.6
                bosluk = r["top"] - onceki["top"] > aralik * 1.55
                kisa_son = onceki["text"].endswith(END_PUNCT) and onceki["x1"] < sag - govde * 2.5
                madde = re.match(r"^\d{1,3}(?:[.)]|\s?[—–-])\s", r["text"])  # "3) ", "4 — " (0.5.13)
                yeni = girinti or bosluk or kisa_son or bool(madde)
        if yeni and cur:
            out.append(cur)
            cur = []
        cur.append(("b" if baslik else "p", r))
        onceki = r
    if cur:
        out.append(cur)
    paras = []
    for grup in out:
        tur = grup[0][0]
        if bas_ayri and tur == "b":  # başlık bloğu satır satır: ilki "b", devamı "b+"
            for n, (_, r) in enumerate(grup):
                paras.append(("b" if n == 0 else "b+", r["text"], r["h"]))
            continue
        metin = TS.join_lines("\n".join(r["text"] for _, r in grup))
        paras.append((tur, metin, max(r["h"] for _, r in grup)))
    return paras


def _dipnot_ayir(rows, h, govde):
    """Önce Stüdyo'nun punto kuralı; bulunamazsa: alt yarıda büyük boşluktan sonra rakamla başlayan satırlar."""
    if not rows:
        return [], []
    # 0.5.15: "⁴ -Allah teâlânın…", "5-Eski…", "7 - Arş'ın…" (çizginin altında, numara üst simge ya da tireli): sayfanın
    # alt kısmında böyle başlayan ve numarası metinde atıf olarak geçen ilk satır. Üstünde gövdeden küçük puntolu satırlar
    # ve büyük boşluk varsa dipnot bölgesi o boşluktan başlar (önceki sayfadan taşan devam: "bildirilmiştir. (Tahrim, 6)")
    for k in range(1, len(rows)):
        m = _NOT_BASI_TIRE.match(rows[k]["text"])
        if not (m and rows[k]["top"] > h * 0.35):
            continue
        no = m.group(1) or m.group(2)
        once = " ".join(r["text"] for r in rows[:k])
        if not (("\ue000" + no + "\ue001") in once or re.search(r"[^\W\d_][.,;:!?\"”»’')]{0,3}" + no + r"(?!\d)", once)):
            continue
        bas = k
        # 0.5.16: yukarıdaki dipnot devamı: büyük boşluktan sonra, yukarısındaki gövdeden belirgin küçük puntolu blok
        # (tek tek satır yüksekliği değil: Surya'da satır yüksekliği oynar)
        for j in range(k - 1, 0, -1):
            if rows[j]["top"] < h * 0.35:
                break
            if rows[j]["top"] - rows[j - 1]["bot"] > govde * 1.5 and _kucuk_blok(rows[:j], rows[j:k], govde):
                bas = j
                break
        if rows[bas]["top"] - rows[bas - 1]["bot"] > govde * 0.8:
            return rows[:bas], rows[bas:]
    surya = sum(1 for r in rows if r.get("surya")) > len(rows) / 2
    if surya:
        # 0.5.16: Surya sayfasında Stüdyo'nun punto kuralı kullanılmaz (sayfa sonundaki kısa "dirler." satırını dipnot
        # sanıyordu); yalnız önceki sayfadan taşan devamdan oluşan dipnot: alt kısımda büyük boşluktan sonra gelen,
        # gövdeden belirgin küçük puntolu blok
        for k in range(3, len(rows)):
            if rows[k]["top"] > h * 0.45 and rows[k]["top"] - rows[k - 1]["bot"] > govde * 1.5 and \
                    _kucuk_blok(rows[:k], rows[k:], govde):
                return rows[:k], rows[k:]
        ana, dip = rows, []
    else:
        ana, dip = TS._dipnot_ayir(rows, h)
    if dip:
        return ana, dip
    for k in range(1, len(rows)):
        a, b = rows[k - 1], rows[k]
        if b["top"] > h * 0.45 and b["top"] - a["bot"] > govde * 2.5 and re.match(r"^\d{1,3}[\s.)]", b["text"]):
            return rows[:k], rows[k:]
    # "(2) İnsan Sûresi, Âyet: 22": parantezli numara, metinle aynı punto, boşluk küçük (ayırma çizgisi). Şart: aynı
    # numara sayfanın metninde atıf olarak geçmeli ("…bolmaz. (2)"); sayfa altındaki numaralı madde dipnot sanılmasın
    for k in range(1, len(rows)):
        a, b = rows[k - 1], rows[k]
        m = _NOT_BASI.match(b["text"])
        if m and b["top"] > h * 0.55 and b["top"] - a["bot"] > govde * 0.8 and \
                _atif_var(" ".join(r["text"] for r in rows[:k]), m.group(1)):
            return rows[:k], rows[k:]
    return rows, []


def _kucuk_blok(ust, alt, govde):
    """0.5.16: alt satırların ortanca yüksekliği hem üstteki gövde satırlarınınkinden (başlıklar sayılmaz) hem sayfanın
    ortanca satır yüksekliğinden belirgin küçük mü (dipnot puntosu)."""
    ust = [r for r in ust if r["h"] <= govde * 1.25]
    if not ust or not alt:
        return False
    o = lambda rs: sorted(r["h"] for r in rs)[len(rs) // 2]
    return o(alt) < o(ust) * 0.9 and o(alt) < govde * 0.95


_NOT_BASI_TIRE = re.compile(r"^\s*(?:\ue000(\d{1,3})\ue001\s*[-–]?|(\d{1,3})\s?[-–])\s*[\"“«'‘]?[^\W\d]")  # 0.5.15/16
_NOT_BASI = re.compile(r"^\s*[(\[]\s?(\d{1,3}|\*{1,3})\s?[)\]]\s*\S")  # 0.5.14: "(*)" yıldızlı dipnot da


_BENZER_ATIF = re.compile(r"(?<=[^\s(\[])\s?[(\[]\s?([IİlıiL|!Oo]{1,2})\s?[)\]]")  # OCR: "(1)" -> "(İİ)"


def _atif_var(metin, no):
    """Metinde "(no)" atıf olarak (bir kelimenin ya da noktalamanın ardından) geçiyor mu. Rakam OCR'da harfe dönmüş
    olabilir ("(İİ)"): metinde tek bir rakama benzeyen parantezli işaret varsa o da sayılır."""
    if re.search(r"[^\s(\[]\s?[(\[]\s?" + re.escape(str(no)) + r"\s?[)\]]", metin):
        return True
    if "\ue000" + str(no) + "\ue001" in metin:  # Surya üst simgesi (0.5.11)
        return True
    return len(_BENZER_ATIF.findall(metin)) == 1


_SAHTE_UST = re.compile(r"(?<=[^\W\d_][.,;:])\s?[!|?'’”\"°](?=\s|$)")
_SON_ISARET = re.compile(r"(?<=[.!?…])\s+['’‘\"”|°]$")


def _tirnak_esli(once, k):
    """Paragraf sonundaki işaret, önceden açılmış bir tırnağın kapanışı mı (gerçek tırnak)?"""
    if k == "”":
        return once.count("“") > once.count("”")
    if k == '"':
        return once.count('"') % 2 == 1
    if k in "’'":
        return once.count("‘") > once.count("’")
    return False


def _okunamayan_ust_simge(metin, sayfa_notu, baglanan):
    """OCR üst simge numarasını çoğu zaman !, ?, ’ ya da ” okur. Noktadan sonra gelen böyle bir işaret, sayfada
    bağlanmamış tek dipnot varsa ona bağlanır. Tırnak ancak eşi yoksa (paragrafta açılmış tırnak yoksa) işaret sayılır."""
    for m in _SAHTE_UST.finditer(metin):
        k = m.group(0)
        once = metin[:m.start()]
        if k in "”\"" and (once.count("“") > once.count("”") or once.count('"') % 2 == 1):
            continue  # açılmış bir tırnağın kapanışı: gerçek tırnak
        if k in "’'" and once.count("‘") > once.count("’"):
            continue
        gid = next(g for g in sayfa_notu.values() if g not in baglanan)
        baglanan.add(gid)
        return metin[:m.start()] + "{{" + gid + "}}" + metin[m.end():]
    return metin


def _notlari_bol(dip_paras):
    """Dipnot paragrafları -> [(numara ya da None (devam), metin)]."""
    out = []
    for p in dip_paras:
        p = UST.sub(r"\1 ", p)
        for parca in re.split(r"(?<![,;:])\s(?=\d{1,3}[\s.)]+[^\W\d])|\s(?=[(\[]\d{1,3}[)\]]\s*[^\W\d])"
                              r"|\s(?=[(\[]\*{1,3}[)\]])|(?<![,;:])\s(?=\d{1,3}\s?[-–]\s?[^\W\d])", p):
            # 0.5.15: "5-Eski", "4 -Allah"; virgülden sonraki sayı ("(Tahrim, 6) Allah") dipnot başı değildir
            parca = parca.strip()
            if not parca:
                continue
            yildiz = re.match(r"^[(\[](\*{1,3})[)\]]\s*(.*)$", parca, re.S)
            if yildiz:  # 0.5.14: "(*)" -> 901, "(**)" -> 902 (metindeki "(*)" atfı aynı numarayla bağlanır)
                out.append((900 + len(yildiz.group(1)), yildiz.group(2).strip()))
                continue
            m_par = re.match(r"^[(\[](\d{1,3})[)\]]\s*(.*)$", parca, re.S)
            m = m_par or re.match(r"^(\d{1,3})(?:\s?[-–]\s?|[\s.)]+)[-–]?\s?(.*)$", parca, re.S)
            onceki = out[-1][0] if out else None
            # 0.5.13: dipnot içindeki sayı ("zamanımızdan 850 sene evvel") yeni dipnot değildir: çıplak sayı ancak
            # sıradaki numaraysa, parantezli numara önceki numaradan biraz büyükse yeni dipnot sayılır
            if m and (onceki is None or (onceki >= 900 and m_par and int(m.group(1)) == 1) or
                      (onceki < 900 and (int(m.group(1)) == onceki + 1 or
                                         (m_par and onceki < int(m.group(1)) <= onceki + 3)))):
                out.append((int(m.group(1)), m.group(2).strip()))
            elif out:
                out[-1] = (out[-1][0], out[-1][1] + " " + parca)
            else:
                out.append((None, parca))
    return out


# ======================= PDF: basılı içindekilerden fihrist =======================
_TOC_SATIR = re.compile(r"^(?P<t>.*?\S)\s*(?:(?:…\s?|(?:[.·•_]\s?){2,})+[^\d]{0,15}?|\s)\s*(?P<n>\d{1,4})[\s,;'’`]{0,4}$")
_TOC_NOKTALI = re.compile(r"^(?P<t>.*?[^\W\d_].*?)\s*(?:[.…·•_]\s?){3,}")
_TOC_ROMA = re.compile(r"^(?P<t>.*?\S)\s*(?:…\s?|(?:[.·•_\-–]\s?){2,})+\s*(?P<n>[ivxlcdmIVXLCDM]{1,7})\s*$")
_TOC_ATLA = re.compile(r"^(sayfa|sahife|s\.|page|pp?\.)$", re.I)


def _toc_temiz(t):
    """Başlık yazısı: sondaki nokta dizisi / OCR kırıntısı atılır."""
    t = re.sub(r"\s*(?:…|(?:[.·•_]\s?){2,}).*$", "", t)
    return t.strip(" .:-–—·•_…")


def _ayni_satir(rows):
    """Aynı satırdaki parçalar birleşir: OCR katmanı başlığı ve sayfa numarasını ayrı satır (çoğu zaman farklı boyda)
    yazmış olabilir. Ölçüt dikey örtüşme; sağda uzakta duran numaranın önüne '…' konur (nokta dizisi yerine)."""
    gruplar = []
    for r in sorted(rows, key=lambda r: (r["top"] + r["bot"]) / 2):
        g = gruplar[-1] if gruplar else None
        if g:
            ust, alt = max(g["top"], r["top"]), min(g["bot"], r["bot"])
            if alt - ust > 0.5 * min(g["bot"] - g["top"], r["bot"] - r["top"]):
                g["p"].append(r)
                g["top"], g["bot"] = min(g["top"], r["top"]), max(g["bot"], r["bot"])
                continue
        gruplar.append({"p": [r], "top": r["top"], "bot": r["bot"]})
    out = []
    for g in gruplar:
        ps = sorted(g["p"], key=lambda r: r["x0"])
        yazi = ps[0]["text"]
        for a, b in zip(ps, ps[1:]):
            uzak = b["x0"] - a["x1"] > 3 * max(a["h"], 5)
            yazi += (" … " if uzak and TS.PAGE_NUM.match(b["text"].strip()) else " ") + b["text"]
        harfli = [r for r in ps if _harf(r["text"]) >= 2]
        out.append(dict(ps[0], text=yazi, x1=ps[-1]["x1"], top=g["top"], bot=g["bot"],
                        h=harfli[0]["h"] if harfli else ps[0]["h"]))
    return out


_TOC_NO_KARAKTER = "0123456789ivxlcdmIVXLCDM"


def _ocr_icindekiler(yol, i, dil):
    """Taranmış içindekiler sayfası. Tesseract nokta dizili satırları olağan kipte atlar ya da birleştirir; seyrek metin
    kipi (--psm 11) kelimeleri tek tek bulur, nokta dizisi düşük güvenli çöp olarak kalır ve atılır. Okunamayan tek haneli
    sayfa numaraları, satırın sağ ucundan (nokta dizisinin bittiği yer) kesilen parçada yeniden okunur."""
    import fitz
    import pytesseract
    from PIL import Image, ImageOps
    page = fitz.open(yol)[i]
    olcek = 72 / OCR_DPI
    img = ImageOps.autocontrast(Image.open(io.BytesIO(page.get_pixmap(dpi=OCR_DPI).tobytes("png"))).convert("L"))
    try:
        d = pytesseract.image_to_data(img, lang=dil, config="--psm 11", output_type=pytesseract.Output.DICT)
    except Exception:
        d = pytesseract.image_to_data(img, config="--psm 11", output_type=pytesseract.Output.DICT)
    kel = []
    for k, t in enumerate(d["text"]):
        t, guven = (t or "").strip(), float(d["conf"][k])
        if not t or guven < 30 or re.fullmatch(r"[\W_]+", t):
            continue  # nokta dizisinden doğan çöp ("LELE", "LL") güveni 25'in altında kalır
        x, y, w, h = d["left"][k], d["top"][k], d["width"][k], d["height"][k]
        if guven < 92 and _harf(t) >= 2:  # şüpheli kelime (ör. "OÖONSOZ"): tek başına kesilip yeniden okunur
            try:
                e = pytesseract.image_to_data(img.crop((x - h // 2, y - h // 2, x + w + h // 2, y + h + h // 2)), lang=dil,
                                              config="--psm 7", output_type=pytesseract.Output.DICT)
                yeni = [(e["text"][j].strip(), float(e["conf"][j])) for j in range(len(e["text"])) if e["text"][j].strip()]
                if len(yeni) == 1 and yeni[0][1] > guven:
                    t, guven = yeni[0]
            except Exception:
                pass
        if guven >= 50:
            kel.append({"t": t, "x0": x, "x1": x + w, "top": y, "h": h})
    satirlar = []
    for w in sorted(kel, key=lambda w: w["top"] + w["h"] / 2):
        orta = w["top"] + w["h"] / 2
        s = next((s for s in satirlar if abs(s["orta"] - orta) < 0.6 * max(s["hs"] + [w["h"]])), None)
        if s:
            s["k"].append(w); s["hs"].append(w["h"])
        else:
            satirlar.append({"orta": orta, "k": [w], "hs": [w["h"]]})
    no_mu = lambda t: re.fullmatch(r"\d{1,4}|[ivxlcdmIVXLCDM]{1,7}", t)
    for s in satirlar:
        s["k"].sort(key=lambda w: w["x0"])
    ters = ImageOps.invert(img).point(lambda v: 255 if v > 90 else 0)
    rows = []
    for s in satirlar:
        ws = s["k"]
        h = st.median(s["hs"])
        metin = " ".join(w["t"] for w in ws)
        kutu = ters.crop((0, int(s["orta"] - h * 0.7), img.size[0], int(s["orta"] + h * 0.7))).getbbox()
        sag = kutu[2] if kutu else 0  # satırdaki son mürekkep: nokta dizisinin sonundaki numara
        if not (len(ws) > 1 and no_mu(ws[-1]["t"])) and sag > ws[-1]["x1"] + h * 3:
            parca = img.crop((int(sag - h * 1.9), int(s["orta"] - h), int(sag + h * 0.4), int(s["orta"] + h)))
            no = ""
            for kip in ("7", "8", "10"):  # tek satır, tek kelime, tek karakter (Tesseract sürümüne göre biri okur)
                try:
                    no = pytesseract.image_to_string(parca, lang=dil, config="--psm " + kip +
                                                     " -c tessedit_char_whitelist=" + _TOC_NO_KARAKTER).strip()
                except Exception:
                    no = ""
                if re.fullmatch(r"\d{1,4}", no) and int(no) > 0:
                    break
            if re.fullmatch(r"\d{1,4}", no) and int(no) > 0:
                metin += " … " + str(int(no))  # "06": numaranın önündeki nokta sıfır okunmuş
        elif len(ws) > 1 and no_mu(ws[-1]["t"]):
            metin = " ".join(w["t"] for w in ws[:-1]) + " … " + ws[-1]["t"]
        metin = TS.norm(metin)
        top = min(w["top"] for w in ws)
        rows.append({"text": metin, "h": h * olcek, "top": top * olcek, "bot": max(w["top"] + w["h"] for w in ws) * olcek,
                     "x0": ws[0]["x0"] * olcek, "x1": max(w["x1"] for w in ws) * olcek, "n": len(metin.split()),
                     "kalin": False, "blok": 0, "ocr": True})
    rows.sort(key=lambda r: r["top"])
    return rows, page.rect.width, page.rect.height


def _icindekiler_devami(yol, sayfalar, satirlar, toc_sayfalari, kaynak=None):
    """Taranmış içindekiler sayfası nokta dizilerine dayanıklı okumayla (_ocr_icindekiler) yeniden okunur (satirlar
    yerinde güncellenir). Hemen ardındaki başlıksız sayfa, satırlarının çoğu içindekiler
    girdisiyse (en az 3) içindekilerin devamıdır."""
    if not toc_sayfalari:
        return toc_sayfalari
    dil = os.environ.get("OCR_LANG", "tur")

    def oku(i):
        asil, yari = kaynak[i] if kaynak else (i, None)
        if sayfalar[i][3] and yari is None:  # çift sayfanın yarısı PDF'te tek başına yok: mevcut okuma
            rows, w, h = _ocr_icindekiler(yol, asil, dil)
            return _sayfa_no_ve_kenar(rows, h)[1]
        return satirlar[i]
    out = []
    for i in toc_sayfalari:
        if i not in out:
            satirlar[i] = oku(i)
            out.append(i)
        j = i + 1
        while j < len(sayfalar) and j not in toc_sayfalari and j - 1 in out and j - out[0] < 8:
            kalan = oku(j)
            yazili = [r for r in kalan if _harf(r["text"]) >= 2]
            if not yazili or len(icindekiler_girdileri([kalan])) < max(3, 0.6 * len(yazili)):
                break
            satirlar[j] = kalan
            out.append(j)
            j += 1
    return sorted(set(out))


def icindekiler_girdileri(sayfa_satirlari):
    """Basılı içindekiler sayfalarının satırları -> [{'baslik', 'no', 'x0', 'ara'?}].
    Ortadaki numarasız satırlar bölüm başlığıdır ('Birinci Bölüm' + 'Kavram ve Terim' -> 'Birinci Bölüm: Kavram ve
    Terim', seviye 1). Soldaki numarasız 'BİRİNCİ BÖLÜM' ardındaki girdiye önek olur; küçük harfle süren satır bölünmüş
    başlığın devamıdır; kalan numarasız satır kendi başına girdidir (yeri komşularından bulunur)."""
    satirlar = []
    for rows in sayfa_satirlari:
        for r in _ayni_satir(rows):
            t = UST.sub(r" \1", r["text"]).strip()  # küçük puntolu numara dipnot işareti sanılmış olabilir
            if not t or _icindekiler_basligi(t) or _TOC_ATLA.match(t):
                continue
            if TS.PAGE_NUM.match(t):  # tek başına numara (aynı hizaya düşmemiş): hemen üstteki numarasız satırın
                son = satirlar[-1] if satirlar else None
                if son and son["no"] is None and not son.get("noktali") and r["top"] - son["bot"] < son["h"]:
                    son["no"] = re.sub(r"\D", "", t) or t.strip().lower()
                continue
            m = _TOC_SATIR.match(_rakam(t)) or _TOC_ROMA.match(t)
            if m and _harf(m.group("t")) >= 2:
                satirlar.append({"t": _toc_temiz(m.group("t")), "no": m.group("n").lower(), "x0": r["x0"],
                                 "x1": r["x1"], "h": r["h"], "bot": r["bot"]})
            elif _TOC_NOKTALI.match(t):  # nokta dizili ama numarası okunamamış: yeri komşularından bulunur
                satirlar.append({"t": _toc_temiz(_TOC_NOKTALI.match(t).group("t")), "no": None, "noktali": True,
                                 "x0": r["x0"], "x1": r["x1"], "h": r["h"], "bot": r["bot"]})
            elif _harf(t) >= 2:
                satirlar.append({"t": _toc_temiz(t), "no": None, "x0": r["x0"], "x1": r["x1"], "h": r["h"], "bot": r["bot"]})
    for s in satirlar:  # "Mantık nedir 9": sondaki sayı başlığın parçası değil
        m = re.match(r"^(.*[^\W\d_].*?)\s+(\d{1,4})$", s["t"])
        if m:
            s["t"] = m.group(1).strip(" .:,;")
            if s["no"] is None and not s.get("noktali"):
                s["no"] = m.group(2)
    numarali = [s for s in satirlar if s["no"]]
    sol = st.median(s["x0"] for s in numarali) if numarali else min((s["x0"] for s in satirlar), default=0)
    sag = st.median(s["x1"] for s in numarali) if numarali else max((s["x1"] for s in satirlar), default=0)
    for s in satirlar:  # ortalanmış numarasız satır: bölüm başlığı
        s["orta"] = s["no"] is None and not s.get("noktali") and s["x0"] > sol + 3 * max(s["h"], 8)
    girdiler, k = [], 0
    while k < len(satirlar):
        s = satirlar[k]
        if s["orta"]:
            parca = [s["t"]]
            while k + 1 < len(satirlar) and satirlar[k + 1]["orta"]:
                k += 1
                parca.append(satirlar[k]["t"])
            if len(parca) > 1 and (BOLUM_NO.match(parca[0]) or _kalip_bolum(parca[0])):
                baslik = parca[0].rstrip(" .:") + ": " + " ".join(parca[1:])
            else:
                baslik = " ".join(parca)
            girdiler.append({"baslik": baslik, "no": None, "x0": s["x0"], "ara": True})
        elif s["no"] is None and not s.get("noktali") and k + 1 < len(satirlar) and not satirlar[k + 1]["orta"]:
            sonraki = satirlar[k + 1]
            if BOLUM_NO.match(s["t"]) or _kalip_bolum(s["t"]):
                sonraki["t"] = s["t"].rstrip(" .:") + ": " + sonraki["t"]
            elif sonraki["t"][:1].islower() or s["x1"] > sol + 0.75 * (sag - sol):  # iki satıra bölünmüş başlık
                sonraki["t"] = s["t"] + " " + sonraki["t"]
                sonraki["x0"] = min(sonraki["x0"], s["x0"])
            elif _buyuk_harfli(s["t"]) and not _buyuk_harfli(sonraki["t"]):
                girdiler.append({"baslik": s["t"], "no": None, "x0": s["x0"], "ara": True})
            else:
                girdiler.append({"baslik": s["t"], "no": None, "x0": s["x0"]})
        else:
            girdiler.append({"baslik": s["t"], "no": s["no"], "x0": s["x0"]})
        k += 1
    girdiler = [g for g in girdiler if _harf(g["baslik"]) >= 2]
    # sıradan sapan numara (OCR hatası: "Mantık nedir 9" ardından "Tarihsel bilgi 5"): yok sayılır, yeri komşularından
    arap = [(k, int(g["no"])) for k, g in enumerate(girdiler) if g["no"] and g["no"].isdigit()]
    for n, (k, no) in enumerate(arap):
        sonraki = arap[n + 1][1] if n + 1 < len(arap) else None
        onceki = arap[n - 1][1] if n > 0 else None
        if sonraki is not None and no > sonraki and (onceki is None or onceki <= sonraki):
            girdiler[k]["no"] = None
    return girdiler


def _sikistir_bosluklu(t):
    from .katalog import sade
    return sade(SAYFA_ISARET.sub("", t))


_NUMARA_ONEKI = re.compile(r"^(?:[^\w\s]+|\d{1,3}[.)\-—–]*|[ivxlcı]{1,5}[.)\-—–]+|[a-zçğıöşü][.)\-—–]+|\S{1,3}<)\s*")


def _kelimeler(t):
    t = _aralik_topla(SAYFA_ISARET.sub("", t).strip()).replace("I", "ı").replace("İ", "i").lower()
    for _ in range(3):  # "il< I. Mantık Nedir?" -> "mantık nedir" (numaralar sadeleştirmeden önce: nokta gerekli)
        yeni = _NUMARA_ONEKI.sub("", t)
        if yeni == t or not yeni:
            break
        t = yeni
    return re.findall(r"\w+", _sikistir_bosluklu(t))


def _toc_puan(metin, baslik):
    """İçindekiler girdisi ile paragraf: 3 aynı, 2 kelime sınırında önek/içerme ("Kıyas" ≠ "Kıyasın tanımı"),
    1 OCR hatalarına dayanıklı benzerlik, 0 değil."""
    a, b = _kelimeler(metin), _kelimeler(baslik)
    if not a or not b:
        return 0
    sa, sb = "".join(a), "".join(b)
    if sa == sb:
        return 3
    if len(sa) <= len(sb) * 1.3 + 12 and (a[:len(b)] == b or b[:len(a)] == a or (len(sb) > 8 and sb in sa)):
        return 2
    return 1 if _bulanik_toc(metin, baslik) else 0


def _sirali_ata(paras, basliklar):
    """Bir sayfanın içindekiler girdileri -> {girdi sırası: paragraf nesnesi}. Girdiler paragraflarla sırayla eşleşir:
    bir girdi, kendisinden sonraki girdinin en iyi eşleştiği paragrafı ve ötesini alamaz ("Kıyas" girdisi, sonraki
    "Kıyasın tanımı"nın aşağısındaki "Kıyas çeşitleri:" satırına oturmasın)."""
    puan = [[_toc_puan(m, b) if t != "sil" else 0 for t, m, _ in paras] for b in basliklar]
    bagimsiz = [max(range(len(paras)), key=lambda k: (p[k], -k)) if paras and max(p) > 0 else None for p in puan]
    out, son = {}, -1
    for n, b in enumerate(basliklar):
        ust = min((bagimsiz[j] for j in range(n + 1, len(basliklar)) if bagimsiz[j] is not None and bagimsiz[j] > son),
                  default=len(paras))
        aday = [k for k in range(son + 1, len(paras)) if puan[n][k] > 0 and
                (k < ust or (k == ust and puan[n][k] > max(puan[j][k] for j in range(n + 1, len(basliklar)))))]
        if not aday and ":" in b:  # "Üçüncü Bölüm: Hüküm ve Önerme": metinde yalnız ikinci kısım
            ikinci = b.split(":", 1)[1]
            aday = [k for k in range(son + 1, min(ust, len(paras))) if _toc_puan(paras[k][1], ikinci) > 0]
        if aday:
            k = max(aday, key=lambda k: (puan[n][k], -k)) if puan[n] and any(puan[n][k] for k in aday) else aday[0]
            out[n], son = paras[k], k
    return out


def _en_iyi(paras, baslik, atla=()):
    """Paragraflar içinde girdiye en çok benzeyen (ilk en yüksek puanlı) sıra ve puanı."""
    en, puan = None, 0
    for k, (t, m, b) in enumerate(paras):
        if k in atla or t == "sil":
            continue
        p = _toc_puan(m, baslik)
        if p > puan:
            en, puan = k, p
            if p == 3:
                break
    return en, puan


def _benzer_toc(metin, baslik):
    """İçindekiler başlığı bu paragraf mı: _benzer_baslik ya da OCR hatalarına dayanıklı benzerlik."""
    return _benzer_baslik(metin, baslik) or _bulanik_toc(metin, baslik)


def _bulanik_toc(metin, baslik):
    """OCR hatalarına dayanıklı benzerlik (%82); sıra sayıları/rakamlar aynı olmalı."""
    import difflib
    a, b = _sikistir(SAYFA_ISARET.sub("", metin)), _sikistir(baslik)
    if len(b) < 8 or not a or len(a) > len(b) * 1.3 + 12:
        return False
    ka, kb = [w for w in _kelimeler(metin) if len(w) > 1], [w for w in _kelimeler(baslik) if len(w) > 1]
    if len(ka) != len(kb) and not (len(kb) >= 3 and abs(len(ka) - len(kb)) <= 1):
        return False  # OCR harf hatası kelime sayısını değiştirmez (bölünmüş kelime: uzun başlıkta bir fark)
    sira = lambda t: re.findall(r"birinci|ikinci|ucuncu|dorduncu|besinci|altinci|yedinci|sekizinci|dokuzuncu|onuncu|\d+|"
                                r"\b[ivxlc]{1,6}\b", _sikistir_bosluklu(t))
    if sira(metin) and sira(baslik) and sira(metin) != sira(baslik):
        return False  # "İKİNCİ BÖLÜM" ≠ "BİRİNCİ BÖLÜM" (yazıca %87 benzer)
    return difflib.SequenceMatcher(None, a[:len(b) + 4], b).ratio() >= 0.82 or \
        difflib.SequenceMatcher(None, a[-len(b) - 4:], b).ratio() >= 0.82


def _yazim_uygun(t):
    return bool(t) and not re.search(r"[^\w\s.,:;?!'’()\-–—]", t)


def _aralik_topla(t):
    """Harf aralıklı başlık ("G İ R İ Ş") -> "GİRİŞ"."""
    kel = t.split()
    if len(kel) >= 3 and all(len(w) <= 2 for w in kel):
        t = "".join(kel)
        if sum(c.isupper() for c in t) > len(t) / 2:
            t = _tr_buyuk(t)
    return t


def _gecersiz_kelime(t):
    try:
        return sum(1 for w in re.findall(r"[^\W\d_]{3,}|\w*\d\w*", t)
                   if re.search(r"\d", w) or not _gecerli(_tr_kucuk(w)))
    except Exception:
        return 0


def _yazim_sec(govde, toc, puan=3):
    """Başlık yazısı: kitabın gövdedeki yazımı (orijinale sadakat) ya da içindekilerdeki. Gövde büyük harfliyse ve
    içindekiler değilse içindekiler (küçük harf OCR'ı çok daha isabetli: "TUMEVAR1M" / "Tümevarım"). Gövdede
    içindekilerde olmayan rakam (OCR: "tanım 1:") ya da tuhaf işaret varsa içindekiler."""
    govde = _aralik_topla(govde)
    if not _yazim_uygun(govde) or set(re.findall(r"\d", " ".join(_kelimeler(govde)))) - set(re.findall(r"\d", toc)):
        return toc
    g, t = _gecersiz_kelime(govde), _gecersiz_kelime(toc)  # Zemberek: "Kıyasm" / "özelliğı" geçersiz
    if g < t and _buyuk_harfli(govde) and not _buyuk_harfli(toc):  # "KAVRAMIN ÖZELLİĞİ" + "Kavramın özelliğı"
        gk, tk = govde.split(), toc.split()
        if len(gk) == len(tk):
            duz = []
            for tw, gw in zip(tk, gk):
                if _gecersiz_kelime(tw) and not _gecersiz_kelime(gw):
                    gw = _tr_kucuk(gw)
                    tw = (_tr_buyuk(gw[0]) + gw[1:]) if tw[:1].isupper() else gw
                duz.append(tw)
            return " ".join(duz)
    if g != t:
        return govde if g < t else toc
    if puan <= 1 or (_buyuk_harfli(govde) and not _buyuk_harfli(toc)):
        return toc
    return govde


def _toc_basligi_yerlestir(paras, k, meta_b, bas):
    """İçindekiler girdisi paras[k]'da bulundu. Birleşmiş başlık satırlarından artan ayrı paragraf olur; 'Birinci Bölüm:
    Kavram ve Terim' girdisinin ikinci kısmı sonraki paragrafsa ona katılır. Döndürür: başlık yazısı; metindeki yazım
    temizse kitabın kendi yazımı (orijinale sadakat), değilse içindekilerdeki."""
    import difflib
    yazi = lambda kk: SAYFA_ISARET.sub("", paras[kk][1]).strip()
    benzer = lambda x, y: difflib.SequenceMatcher(None, _sikistir(x), _sikistir(y)).ratio()
    bas_kismi, artan = _basligi_ayir(paras[k][1], bas)
    if artan:  # "İKİNCİ BÖLÜM ÖNERMELER Önermenin Tanımı" -> başlık + ayrı paragraf
        paras[k:k + 1] = [("p", bas_kismi, paras[k][2]), ("p", artan, paras[k][2])]
        for kk in sorted([kk for kk in meta_b if kk > k], reverse=True):
            meta_b[kk + 1] = meta_b.pop(kk)
    govde = yazi(k)
    hedef, kk = _sikistir(bas), k + 1  # iki satıra bölünmüş başlık: devamı sonraki paragraf(lar)da
    while ":" not in bas and hedef.startswith(_sikistir(govde)) and len(_sikistir(govde)) < len(hedef) and kk < len(paras) \
            and kk not in meta_b and paras[kk][0] != "sil":
        parca = yazi(kk)
        if not parca or not hedef[len(_sikistir(govde)):].startswith(_sikistir(parca)):
            break
        paras[kk] = ("sil", "", 0)
        govde += " " + parca
        kk += 1
    if ":" in bas:
        on, arka = [x.strip() for x in bas.split(":", 1)]
        if benzer(govde, on) >= 0.85 and k + 1 < len(paras) and k + 1 not in meta_b and paras[k + 1][0] != "sil":
            ikinci = yazi(k + 1)
            if benzer(ikinci, arka) >= 0.75 and len(_sikistir(ikinci)) <= len(_sikistir(arka)) + 6:
                paras[k + 1] = ("sil", "", 0)
                govde, ikinci = _aralik_topla(govde), _aralik_topla(ikinci)
                if _yazim_uygun(govde) and _yazim_uygun(ikinci):
                    return govde.rstrip(" .:") + ": " + ikinci  # "BİRİNCİ BÖLÜM: KAVRAM VE TERİM" (büyük puntolu, isabetli)
        return bas
    puan = _toc_puan(govde, bas)
    if puan >= 1 and len("".join(_kelimeler(govde))) <= len("".join(_kelimeler(bas))) + 3:
        return _yazim_sec(govde, bas, puan)
    return bas


def _kisa_parcalari_birlestir(rows):
    """Aynı satırda yan yana duran kısa parça (büyük ilk harf "G" + "İ R İş") bir sonraki parçayla birleşir."""
    out = []
    for r in rows:
        p = out[-1] if out else None
        if p and len(p["text"].strip()) <= 3 and r["x0"] >= p["x1"] - 2 and r["x0"] - p["x1"] < 2 * max(p["h"], r["h"]) \
                and min(p["bot"], r["bot"]) - max(p["top"], r["top"]) > 0.4 * min(p["bot"] - p["top"], r["bot"] - r["top"]):
            out[-1] = dict(r, text=p["text"].strip() + " " + r["text"], x0=p["x0"], h=max(p["h"], r["h"]),
                           top=min(p["top"], r["top"]), bot=max(p["bot"], r["bot"]))
            continue
        out.append(r)
    return out


def _devam_satirlarini_birlestir(paras, meta_b, asil_b=()):
    """İçindekiler kipinde başlık blokları satır satır ayrılır; başlık olarak kullanılmayan devam satırları ("p+")
    önceki satırla yeniden tek paragraf olur."""
    out, yeni_meta = [], {}
    for k, (t, m, b) in enumerate(paras):
        if t == "p+" and k not in meta_b and out and out[-1][0] == "p" and (len(out) - 1) not in yeni_meta:
            out[-1] = ("p", TS.join_lines(out[-1][1] + "\n" + m), max(out[-1][2], b))
            continue
        # başlık sanılan (kalın) satırla başlayıp alt satırda küçük harfle süren cümle: tek paragraf
        if t == "p" and k not in meta_b and k - 1 in asil_b and k - 1 not in meta_b and out and out[-1][0] == "p" \
                and (len(out) - 1) not in yeni_meta and SAYFA_ISARET.sub("", m).lstrip()[:1].islower() \
                and not SAYFA_ISARET.sub("", out[-1][1]).rstrip().endswith(END_PUNCT):
            out[-1] = ("p", out[-1][1] + " " + m, max(out[-1][2], b))
            continue
        if k in meta_b:
            yeni_meta[len(out)] = meta_b[k]
        out.append(("p" if t == "p+" else t, m, b))
    return out, yeni_meta


def _basligi_ayir(metin, baslik):
    """Paragraf başlıkla başlıyor ve devam ediyorsa (birleşmiş başlık satırları): (başlık kısmı, artan) ; değilse (metin, "")."""
    hedef = _sikistir(baslik)
    kel = metin.split(" ")
    for n in range(1, len(kel)):
        on = _sikistir(SAYFA_ISARET.sub("", " ".join(kel[:n])))
        if on == hedef:
            artan = " ".join(kel[n:]).strip()
            return (" ".join(kel[:n]), artan) if _harf(artan) >= 2 else (metin, "")
        if len(on) > len(hedef) or not hedef.startswith(on):
            break
    return metin, ""


def _kalip_bolum(t):
    return bool(re.match(r"^(bölüm|kısım|fasıl|bab|kitap|makale)\s+([ivxlc]{1,6}|\d{1,3})[.:]?$", t, re.I))


def _toc_seviyeleri(girdiler):
    """Seviye (en çok 3). Ortadaki bölüm başlıkları ('ara') 1; öteki girdiler girinti kümelerine göre, bölüm başlığı
    varsa onun bir altı. Girinti yoksa: büyük harfli / 'Birinci Bölüm' kalıplı girdiler üst seviye.
    İlk bölüm başlığından önceki girdiler (Önsöz gibi) 1."""
    duz = [g for g in girdiler if not g.get("ara")]
    kumeler = []
    for x in sorted({round(g["x0"]) for g in duz}):
        if not kumeler or x - kumeler[-1][-1] > 8:
            kumeler.append([x])
        else:
            kumeler[-1].append(x)
    kume = lambda g: next(k for k, c in enumerate(kumeler) if round(g["x0"]) in c)
    say = collections.Counter(kume(g) for g in duz)
    girintili = len([k for k in say if say[k] >= 2]) >= 2
    ara_var = any(g.get("ara") for g in girdiler)
    ust = [g for g in duz if _buyuk_harfli(g["baslik"]) or BOLUM_NO.match(g["baslik"].split(":")[0])]
    ilk_ara = next((k for k, g in enumerate(girdiler) if g.get("ara")), len(girdiler))
    for k, g in enumerate(girdiler):
        if g.get("ara") or (ara_var and k < ilk_ara):
            g["seviye"] = 1
        elif girintili:
            g["seviye"] = kume(g) + 1 + ara_var
        elif ara_var:
            g["seviye"] = 2
        else:
            g["seviye"] = 1 if (g in ust or len(ust) == len(duz)) else 2
    kullanilan = sorted({g["seviye"] for g in girdiler})
    for g in girdiler:
        g["seviye"] = min(kullanilan.index(g["seviye"]) + 1, 3)


def icindekiler_fihristi(girdiler, nolar, kalan, paras_of):
    """Basılı içindekiler girdileri -> {pdf_sayfa_sırası: [(seviye, başlık)]} ya da None (güvenilir değil).
    Başlık, gösterdiği sayfada ve yakınında (-1..+2, sıra bozulmadan) metinde aranır; bulunamazsa gösterdiği sayfanın başına.
    paras_of(i): i. PDF sayfasının paragrafları [(tür, metin, boy)]."""
    if len(girdiler) < 3:
        return None
    sayfa_of, kalan_set = {}, set(kalan)
    for i, no in enumerate(nolar):  # yalnız metin sayfaları (içindekiler sayfasına yanlış okunmuş numara düşebilir)
        if no and no not in sayfa_of and i in kalan_set:
            sayfa_of[no] = i
    eslenen = [g for g in girdiler if g["no"] in sayfa_of]
    if len(eslenen) < max(3, 0.6 * sum(1 for g in girdiler if g["no"])):
        return None  # numaraların çoğu kitapta yok
    arap = [int(g["no"]) for g in eslenen if g["no"].isdigit()]
    eslenen = [g for g in girdiler if g["no"] in sayfa_of or g["no"] is None]
    if arap and sum(1 for a, b in zip(arap, arap[1:]) if b >= a) < 0.8 * (len(arap) - 1):
        return None  # numaralar artmıyor: içindekiler değil (ör. dizin, kronoloji)
    _toc_seviyeleri(girdiler)
    imler, bulunan, onceki, bekleyen = collections.defaultdict(list), 0, -1, []
    for k, g in enumerate(eslenen):
        if g["no"] is None:
            sonraki = next((sayfa_of[x["no"]] for x in eslenen[k + 1:] if x["no"]), max(kalan_set))
            adaylar = [i for i in sorted(kalan_set) if max(onceki, 0) <= i <= sonraki]
        else:
            hedef = sayfa_of[g["no"]]
            adaylar = sorted((i for i in range(hedef - 1, hedef + 3) if i in kalan_set and i >= onceki),
                             key=lambda i: (abs(i - hedef), i))
        puanlar = [(_en_iyi(paras_of(i), g["baslik"])[1], -n, i) for n, i in enumerate(adaylar)]
        en = max(puanlar, default=(0, 0, None))
        yer = en[2] if en[0] > 0 else None
        if yer is not None:
            bulunan += 1
            onceki = yer  # sıra yalnız metinde bulunan başlıklarla ilerler (yanlış numara zincirleme bozmasın)
            for yedek, gg in bekleyen:  # bulunamayanlar: kendi sayfaları bundan önceyse orada, değilse bunun önünde
                imler[yedek if yedek is not None and yedek < yer else yer].append((gg["seviye"], gg["baslik"]))
            bekleyen = []
        elif g["no"] is None:
            if g.get("ara") and adaylar:  # bulunamayan bölüm başlığı: sonraki bulunan girdinin önünde
                bekleyen.append((None, g))
            continue
        else:
            yedek = next((i for i in sorted(kalan_set) if i >= max(hedef, onceki)), None)
            if yedek is not None and yedek - hedef <= 2:
                bekleyen.append((yedek, g))
            continue
        imler[yer].append((g["seviye"], g["baslik"]))
    for yedek, gg in bekleyen:  # sonda kalanlar
        if yedek is not None:
            imler[yedek].append((gg["seviye"], gg["baslik"]))
    if bulunan < max(2, 0.4 * len(eslenen)):
        return None  # başlıkların çoğu metinde yok: sayfa numaraları başka baskıya ait olabilir
    return imler


def pdf_oku(yol, ilerleme=None):
    """PDF -> (öğeler, dipnotlar, bilgi). öğe: {'tur': 'baslik'|'p', 'metin', 'boy'} ; metinde sayfa/dipnot işaretleri."""
    _gerekli()
    sayfalar, bilgi = pdf_sayfalari(yol, ilerleme)
    if ilerleme:
        ilerleme("Sayfa yapısı çıkarılıyor")
    n = len(sayfalar)
    nolar, satirlar = [], []
    for rows, w, h, _ in sayfalar:
        no, kalan = _sayfa_no_ve_kenar(rows, h)
        nolar.append(no)
        satirlar.append(kalan)
    nolar = _eksik_numaralari_doldur(nolar)
    satirlar = _kenar_numarasi_dizisi(satirlar)
    bilgi["kapak_baslik"] = _kapak_basligi(satirlar[:3])
    dolu = next((rows for rows in satirlar[:3] if any(_harf(r["text"]) >= 3 for r in rows)), [])
    bilgi["kapak_satirlari"] = [(r["text"], r["h"], r["top"]) for r in sorted(dolu, key=lambda r: r["top"])
                                if 2 <= len(r["text"]) <= 80 and _harf(r["text"]) >= 2]
    # ön ve son sayfalar (kapak, künye, içindekiler): Stüdyo'nun kuralı
    tut = TS.on_ve_son_sayfalari_at([[str(i)] + [r["text"] for r in satirlar[i]] for i in range(n)])
    kalan = [int(s[0]) for s in tut]
    # basılı içindekiler sayfaları (başta, sonda ya da önsözden sonra): metinden çıkar, kitabın sonuna eklenir
    toc_sayfalari = [i for i in range(n) if (i < max(20, n // 5) or i >= n - 15) and _icindekiler_sayfasi(satirlar[i])]
    toc_sayfalari = _icindekiler_devami(yol, sayfalar, satirlar, toc_sayfalari, bilgi.get("kaynak_sayfa"))
    if toc_sayfalari:  # içindekiler sayfalarının arasına/devamına düşen numaralı satır sayfaları da
        kalan = [i for i in kalan if i not in toc_sayfalari]
    satirlar_k = _tekrar_edenleri_at([satirlar[i] for i in kalan])
    govde_k, govde_o = _govde_boyu(satirlar_k, False), _govde_boyu(satirlar_k, True)
    govde = govde_k or govde_o or 11
    tum = [r for rows in satirlar_k for r in rows]
    kalin_oran = sum(len(r["text"]) for r in tum if r["kalin"]) / max(1, sum(len(r["text"]) for r in tum))
    ogeler, notlar, not_sayac = [], {}, [0]
    son_not = None
    # PDF'in kendi yer imleri varsa başlıklar oradan (bizim başlık tanımamız devre dışı)
    imler = collections.defaultdict(list)
    tum_imler = bilgi.get("yer_imleri") or []
    anlamli = [x for x in tum_imler if not re.fullmatch(r"(page|sayfa|pg|p|s)?\.?\s*\d{1,4}|[ivxlcdm]{1,6}", x[1].strip(), re.I)]
    if len(anlamli) < 0.5 * len(tum_imler):
        anlamli = []  # yer imlerinin çoğu "Page 1, Page 2…": tarama programının koyduğu, fihrist değil
    for sv, bas, sy in anlamli:
        if sy - 1 in set(kalan):
            imler[sy - 1].append((min(sv, 3), bas))
    yer_imi_modu = sum(len(v) for v in imler.values()) >= 3
    parcalar = []  # her kalan sayfa: (gövde boyu, dipnot satırları, paragraflar)
    for j, i in enumerate(kalan):
        g = (govde_o if sayfalar[i][3] else govde_k) or govde
        ana, dip = _dipnot_ayir(satirlar_k[j], sayfalar[i][2], g)
        parcalar.append((g, dip, _sayfa_paragraflari(ana, sayfalar[i][1], g, kalin_oran)))
    icindekiler_modu = False
    if not yer_imi_modu and toc_sayfalari:  # yer imi yoksa: basılı içindekiler sayfası fihrist kaynağı
        sira = {i: j for j, i in enumerate(kalan)}
        ayri = []  # başlık satırları ayrı paragraf ("GİRİŞ" + "I. Mantık Nedir?" tek blok olmasın)
        for j, i in enumerate(kalan):
            g, dip, _ = parcalar[j]
            ana, _ = _dipnot_ayir(satirlar_k[j], sayfalar[i][2], g)
            ana = _kisa_parcalari_birlestir(ana)
            ayri.append((g, dip, _sayfa_paragraflari(ana, sayfalar[i][1], g, kalin_oran, bas_ayri=True)))
        toc_imler = icindekiler_fihristi(icindekiler_girdileri([satirlar[i] for i in toc_sayfalari]), nolar, kalan,
                                         lambda i: ayri[sira[i]][2])
        if toc_imler:
            imler, yer_imi_modu, icindekiler_modu, parcalar = toc_imler, True, True, ayri
            toc_girdi = sum(len(v) for v in imler.values())
    if yer_imi_modu:
        bilgi["yapi"] = "fihrist"
        bilgi["fihrist_kaynagi"] = "icindekiler" if icindekiler_modu else "yer_imleri"
    son_seviye, toc_girdi = 1, locals().get("toc_girdi", 0)
    for j, i in enumerate(kalan):
        g, dip, paras = parcalar[j]
        dip_paras = [TS.join_lines(r["text"]) for r in dip]
        # sayfanın dipnotları: numaralı olanlar yeni, numarasız baştaki parça önceki sayfanın notunun devamı
        sayfa_notu = {}
        for no, metin in _notlari_bol(dip_paras):
            if no is None:
                if son_not:
                    notlar[son_not] += " " + metin
                continue
            not_sayac[0] += 1
            gid = "n%04d" % not_sayac[0]
            notlar[gid] = metin
            sayfa_notu.setdefault(no, gid)
            son_not = gid
        baglanan = set()

        def bagla(m, sayfa_notu=sayfa_notu, baglanan=baglanan):
            gid = sayfa_notu.get(int(m.group(1)))
            if gid and gid not in baglanan:
                baglanan.add(gid)
                return "{{" + gid + "}}"
            return ""
        yeni_paras = []
        for tur, metin, boy in paras:
            metin = UST.sub(bagla, metin)
            if tur == "p" and any(n > 900 for n in sayfa_notu):  # 0.5.14: "…yazılmamıştır.(*)"
                metin = re.sub(r"(?<=[^\s(\[])\s?[(\[]\s?(\*{1,3})\s?[)\]]",
                               lambda m: bagla(re.match(r"(\d+)", str(900 + len(m.group(1))))), metin)
            if tur == "p" and sayfa_notu:  # okunamamış üst simge: kelimeye yapışık rakam, noktadan sonra "!"
                metin = YAPISIK.sub(bagla, metin)
                # parantezli atıf "…bolmaz. (2)" (bir kelimenin/noktalamanın ardından; paragraf başındaki madde no değil)
                metin = re.sub(r"(?<=[^\s(\[])\s?[(\[]\s?(\d{1,3})\s?[)\]]",
                               lambda m: bagla(m) if int(m.group(1)) in sayfa_notu and
                               sayfa_notu[int(m.group(1))] not in baglanan else m.group(0), metin)
                # OCR katmanı üst simgeyi normal boyda yazmış: "edilmesine 74 niyet" (sayfada 74 numaralı dipnot varsa)
                metin = re.sub(r"(?:(?<=[^\W\d_])|(?<=[^\W\d_][,;:])) (\d{1,3})(?= [^\W\d_])",
                               lambda m: (" " + bagla(m)) if int(m.group(1)) in sayfa_notu and
                               sayfa_notu[int(m.group(1))] not in baglanan else m.group(0), metin)
                if sayfalar[i][3] and len(sayfa_notu) - len(baglanan) == 1:
                    metin = _okunamayan_ust_simge(metin, sayfa_notu, baglanan)
                bos = [n for n, g_ in sayfa_notu.items() if g_ not in baglanan]
                if len(bos) == 1 and len(_BENZER_ATIF.findall(metin)) == 1:  # OCR: "(1)" -> "(İİ)"; tek aday, tek not
                    metin = _BENZER_ATIF.sub(lambda m: bagla(re.match(r"(\d+)", str(bos[0]))), metin)
            elif tur == "p":
                metin = YAPISIK.sub("", metin)
            if tur == "p" and sayfalar[i][3]:  # bağlanamayan sahipsiz son işaret: sil
                son = _SON_ISARET.search(metin)
                if son and not _tirnak_esli(metin[:son.start()], son.group(0).strip()):
                    metin = metin[:son.start()]
            yeni_paras.append((tur, metin, boy))
        # metinde atfı bulunamayan notlar: sayfanın son paragrafının sonuna bağlanır (not kaybolmasın)
        bosta = [g for g in sayfa_notu.values() if g not in baglanan]
        if bosta:
            for k in range(len(yeni_paras) - 1, -1, -1):
                if yeni_paras[k][0] == "p":
                    tur, metin, boy = yeni_paras[k]
                    yeni_paras[k] = (tur, metin + "".join("{{" + g + "}}" for g in bosta), boy)
                    break
            else:
                yeni_paras.append(("p", "".join("{{" + g + "}}" for g in bosta), govde))
        meta_b = {}  # yeni_paras sırası -> (seviye, yer imi başlığı)
        if yer_imi_modu:
            asil_b = {k for k, (t, m, b) in enumerate(yeni_paras) if t in ("b", "b+")}
            # "b+": aynı başlık bloğunun devam satırı (içindekiler kipinde satırlar ayrı); eşleşmezse geri birleşir
            yeni_paras = [("p" if t == "b" else "p+" if t == "b+" else t, m, b) for t, m, b in yeni_paras]
            eklenen = 0
            sayfa_imleri = imler.get(i, [])
            atanan = _sirali_ata(yeni_paras, [b for _, b in sayfa_imleri]) if icindekiler_modu else {}
            for n_im, (sv, bas) in enumerate(sayfa_imleri):
                if icindekiler_modu:
                    hedef_p = atanan.get(n_im)
                    k = next((kk for kk, pp in enumerate(yeni_paras) if pp is hedef_p and kk not in meta_b), None) \
                        if hedef_p is not None else None
                else:
                    k = next((k for k, (t, m, b) in enumerate(yeni_paras) if k not in meta_b and _benzer_baslik(m, bas)), None)
                bulundu = k is not None
                if k is None:  # sayfada bulunamadı: yer imi kipinde sayfanın başına; içindekilerde son başlığın ardına
                    yer = (max(meta_b) + 1) if (icindekiler_modu and meta_b) else eklenen
                    yeni_paras.insert(yer, ("b", bas, 0))
                    meta_b = {(kk + 1 if kk >= yer else kk): v for kk, v in meta_b.items()}
                    asil_b = {(kk + 1 if kk >= yer else kk) for kk in asil_b}
                    k = yer
                    eklenen += 1
                if icindekiler_modu and bulundu:
                    once = len(yeni_paras)
                    bas = _toc_basligi_yerlestir(yeni_paras, k, meta_b, bas)
                    if len(yeni_paras) > once:  # paragraf bölündü: sonraki sıralar bir kaydı
                        asil_b = {(kk + 1 if kk > k else kk) for kk in asil_b}
                yeni_paras[k] = ("b", yeni_paras[k][1], yeni_paras[k][2])
                meta_b[k] = (sv, bas)
            if icindekiler_modu and toc_girdi < 15:  # kısa içindekiler (yalnız bölümler): gövdedeki belirgin büyük
                for k in sorted(asil_b):                    # başlık önceki girdinin altına eklenir
                    t, m, b = yeni_paras[k]
                    if k not in meta_b and t == "p" and b >= g * 1.12 and anlamli_baslik(SAYFA_ISARET.sub("", m)):
                        meta_b[k] = (None, None)
                        yeni_paras[k] = ("b", m, b)
        if icindekiler_modu:
            yeni_paras, meta_b = _devam_satirlarini_birlestir(yeni_paras, meta_b, asil_b)
        etiket = "\ue002" + nolar[i] + "\ue003" if nolar[i] else ""
        for k, (tur, metin, boy) in enumerate(yeni_paras):
            if tur == "sil":
                continue
            ilk = k == 0
            onceki = ogeler[-1] if ogeler else None
            # sayfa geçişinde bölünen paragraf: öncekiyle birleştir
            kuyruk_m = re.search(r"((?:\{\{n\d{4,}\}\})+)$", onceki["metin"]) if onceki else None
            if ilk and tur == "p" and onceki and onceki["tur"] == "p" and metin[:1].islower() and not onceki.get("koru") \
                    and not onceki["metin"][:kuyruk_m.start() if kuyruk_m else None].endswith(END_PUNCT):
                kuyruk = kuyruk_m.group(1) if kuyruk_m else ""
                sol = onceki["metin"][:kuyruk_m.start()] if kuyruk_m else onceki["metin"]
                if sol.endswith(("-", "‐")):
                    kel_sol = re.search(r"([^\W\d_]+)[-‐]$", sol)
                    kel_sag = re.match(r"([^\W\d_]+)", metin)
                    if kel_sol and kel_sag:
                        birlesik = DZ._birlesik(kel_sol.group(1), kel_sag.group(1)) or (kel_sol.group(1) + kel_sag.group(1))
                        onceki["metin"] = sol[:kel_sol.start()] + etiket + birlesik + kuyruk + metin[kel_sag.end():]
                    else:
                        onceki["metin"] = sol[:-1] + etiket + metin + kuyruk
                else:
                    onceki["metin"] = sol + kuyruk + " " + etiket + metin
                continue
            if k in meta_b:  # yer iminden başlık: metni yer imindeki yazı
                sv, bas = meta_b[k]
                if sv is None:  # içindekilerde olmayan başlık: kendi yazısı, bir alt seviye
                    ogeler.append({"tur": "baslik", "metin": (etiket if ilk else "") + metin, "boy": boy,
                                   "seviye": min(son_seviye + 1, 3), "fihrist": True, "onar": True})
                    continue
                son_seviye = sv
                ogeler.append({"tur": "baslik", "metin": (etiket if ilk else "") + bas, "boy": 0, "seviye": sv, "fihrist": True,
                               "onar": icindekiler_modu})
                continue
            ogeler.append({"tur": "baslik" if tur == "b" else "p", "metin": (etiket if ilk else "") + metin, "boy": boy,
                           "ocr": sayfalar[i][3]})
        if not yeni_paras and etiket and ogeler:
            pass  # boş sayfa: numarası atlanır
        if ilerleme and j % 20 == 0:
            ilerleme(f"Sayfa yapısı: {j + 1}/{len(kalan)}")
    # basılı içindekiler/fihrist sayfaları kitaptan silinir (0.5.8: EPUB'da liste ve içindekiler yok)
    if not yer_imi_modu:
        _sahte_basliklari_ayikla(ogeler)
    _kopuk_paragraflari_birlestir(ogeler)
    if not yer_imi_modu and _bolum_iskeleti(ogeler):  # içindekiler/yer imi yok, metinde "… Bölüm" başlıkları var
        bilgi["yapi"] = "fihrist"
        bilgi["fihrist_kaynagi"] = "bolum_basliklari"
    bilgi["govde_boyu"] = govde
    return ogeler, notlar, bilgi


def _sahte_basliklari_ayikla(ogeler):
    """Yazı boyu tahmininin başlık sandığı paragraf parçaları ve tarama lekeleri: küçük harfle başlayan ("rılmıştır.",
    "dür.»" — paragrafın son satırı), 3'ten az harfli ("HI", "BH") ya da "|" ile biten satır başlık değildir; paragraf
    olur (ardından kopuk paragraf birleştirme onu cümlesine geri ekler)."""
    for k, o in enumerate(ogeler):
        if o["tur"] != "baslik" or o.get("koru") or o.get("fihrist"):
            continue
        t = SAYFA_ISARET.sub("", o["metin"]).strip()
        if re.match(r"[a-zçğıöşüâîû]", t) or _harf(t) < 3 or t.endswith("|"):
            ogeler[k] = dict(o, tur="p")


def _kopuk_paragraflari_birlestir(ogeler):
    """Cümlesi yarıda kalan paragrafın devamı ayrı paragraf olmuş (kalın ayet meali satır satır, sayfa geçişi, araya
    giren dipnot): önceki paragraf cümle sonu işaretiyle bitmiyorsa ve sonraki küçük harfle başlıyorsa birleşir
    ("…bu ni'met-" + "leri, siz iyi…"). Satır içinde kalan tire ("kay- bolmaz") sonraki temizlikte birleşir.
    Liste maddeleri ("b) …", "0) …") ve korunan (sona taşınan içindekiler) satırlar birleşmez."""
    yeni = []
    for o in ogeler:
        p = yeni[-1] if yeni else None
        if p and o["tur"] == "p" and p["tur"] == "p" and not o.get("koru") and not p.get("koru"):
            once = re.sub(r"\{\{n\d+\}\}|[\ue000-\ue003]", "", SAYFA_ISARET.sub("", p["metin"])).rstrip()
            sonra = SAYFA_ISARET.sub("", o["metin"]).lstrip()
            if once and _LATIN.search(once) and not re.match(r"[a-zçğıöşü\d][)\].]\s", sonra) and (
                    (not once.endswith(END_PUNCT) and _AR.search(once[-30:])) or
                    (once.endswith((")", "﴾")) and _AR.search(once[-80:]) and (sonra.startswith("«") or sonra[:1].islower()))):
                yeni[-1] = dict(p, metin=p["metin"].rstrip() + " " + o["metin"].lstrip())
                continue
            if once and not once.endswith(END_PUNCT) and re.match(r"[a-zçğıöşüâîû]", sonra) and \
                    not re.match(r"[a-zçğıöşü][)\].]\s", sonra) and not re.match(r"^\(?[a-zçğıöşü\d][)\].]\s", once):
                yeni[-1] = dict(p, metin=p["metin"].rstrip() + " " + o["metin"].lstrip())
                continue
        yeni.append(o)
    ogeler[:] = yeni


_BOLUM_ON = re.compile(r"^(birinci|ikinci|üçüncü|dördüncü|beşinci|altıncı|yedinci|sekizinci|dokuzuncu|onuncu|"
                       r"on\s?birinci|on\s?ikinci|\d{1,2}\.?|[ivxlc]{1,6}\.?)\s+(bölüm|kısım|fasıl|bab|kitap|makale)\b")
_ANA_BASLIK = re.compile(r"^(önsöz|ön söz|giriş|mukaddime|takdim|sunuş|başlarken|sonuç|son söz|hâtime|hatime|"
                         r"netice|bibliyografya|kaynakça|kaynaklar)\b[.:]?$")


def _bolum_iskeleti(ogeler):
    """Basılı içindekiler ve yer imi yoksa fihrist iskeleti metindeki bölüm başlıklarından: "Birinci Bölüm …"
    (ve Önsöz, Giriş, Mukaddime, Sonuç gibi) başlıklar seviye 1, öteki başlıklar seviye 2. İki nokta, ünlem ya da soru
    işaretiyle biten satır ("Hazret-i … buyurur:", "Tarsın!") başlık değildir. Bölüm satırının hemen ardından gelen başlık
    ("İKİNCİ BÖLÜM" + "TEVBE") onunla birleşir. En az iki bölüm başlığı yoksa hiçbir şey yapılmaz."""
    temiz = lambda o: SAYFA_ISARET.sub("", o["metin"]).strip()
    bas = [o for o in ogeler if o["tur"] == "baslik" and not o.get("koru")]
    if sum(1 for o in bas if _BOLUM_ON.match(_tr_kucuk(temiz(o)))) < 2:
        return False
    yeni, onceki_bolum = [], None
    for o in ogeler:
        if o["tur"] != "baslik" or o.get("koru"):
            yeni.append(o)
            onceki_bolum = None
            continue
        t = temiz(o)
        kucuk = _tr_kucuk(t)
        if t.endswith((":", "!", "?")) or len(t.split()) > 14 or not _harf(t):
            yeni.append(dict(o, tur="p"))  # konuşma/alıntı girişi, başlık değil
            onceki_bolum = None
            continue
        if _BOLUM_ON.match(kucuk) or _ANA_BASLIK.match(kucuk):
            o = dict(o, seviye=1, fihrist=True, onar=True)
            yeni.append(o)
            onceki_bolum = o if _BOLUM_ON.match(kucuk) and len(t.split()) <= 3 else None
            continue
        harfler = [c for c in t if c.isalpha()]
        if len(harfler) < 4 or sum(c.isupper() for c in harfler) < 0.8 * len(harfler):
            yeni.append(dict(o, tur="p"))  # alt başlık değil (karışık harfli cümle, tarama lekesi)
            onceki_bolum = None
            continue
        if onceki_bolum is not None:  # "İKİNCİ BÖLÜM" + "TEVBE" -> "İKİNCİ BÖLÜM: TEVBE"
            onceki_bolum["metin"] = onceki_bolum["metin"].rstrip(" .:") + ": " + t
            onceki_bolum = None
            continue
        yeni.append(dict(o, seviye=2, fihrist=True, onar=True))
    ogeler[:] = yeni
    _kopuk_paragraflari_birlestir(ogeler)  # paragrafa dönen parçalar cümlesine
    return True


# ======================= EPUB / DOCX / TXT =======================
_NOT_BASLIK = re.compile(r"^(notlar|dipnotlar|sonnotlar|açıklamalar|notes|endnotes|footnotes)[.:]?$", re.I)
_NOKTALI = re.compile(r"(\.{3,}|…{2,}|(\. ){3,})\s*\d{0,4}\s*$|\s\d{1,4}\s*$")


def _on_temizlik(ogeler):
    """İçindekiler bölümü (noktalı/numaralı satırlar), baştaki künye satırları ve içi boş Notlar başlıkları atılır."""
    yazi = lambda o: SAYFA_ISARET.sub("", o["metin"]).strip()
    out, i, n = [], 0, len(ogeler)
    while i < n:
        o = ogeler[i]
        if o["tur"] == "baslik" and TS.ICINDEKILER.search(yazi(o)) and len(yazi(o)) < 40 and i < max(10, n * 0.3):
            j = i + 1
            while j < n and ogeler[j]["tur"] == "p" and (_NOKTALI.search(yazi(ogeler[j])) or
                                                        TS.KUNYE.search(yazi(ogeler[j])) or len(yazi(ogeler[j])) < 60):
                j += 1
            i = j
            continue
        if o["tur"] == "p" and i < max(10, n * 0.05) and TS.KUNYE.search(yazi(o)) and len(yazi(o)) < 200:
            i += 1
            continue
        if o["tur"] == "baslik" and _NOT_BASLIK.match(yazi(o)) and (i + 1 == n or ogeler[i + 1]["tur"] == "baslik"):
            i += 1
            continue
        out.append(o)
        i += 1
    return out

_BLOK = ["h1", "h2", "h3", "h4", "h5", "h6", "p", "li", "blockquote", "dd", "dt", "pre"]


def epub_oku(yol, ilerleme=None):
    """EPUB -> (öğeler, dipnotlar, bilgi). h1-h6 başlık; noteref bağlantıları dipnot; pagebreak işaretleri sayfa."""
    _gerekli()
    from bs4 import BeautifulSoup, NavigableString
    from ebooklib import ITEM_DOCUMENT, epub
    book = epub.read_epub(yol, options={"ignore_ncx": False})  # EPUB2 (Calibre) fihristi NCX'te
    belgeler = []
    for idref, _ in book.spine:
        item = book.get_item_with_id(idref)
        # bazı EPUB'lar .html sayfalarının türünü yanlış işaretler: uzantıya da bakılır
        if item and (item.get_type() == ITEM_DOCUMENT or item.get_name().lower().endswith((".html", ".htm", ".xhtml"))):
            belgeler.append((os.path.basename(item.get_name()), BeautifulSoup(item.get_content(), "html.parser")))
    # 1) dipnot hedefleri: kısa (rakam, [1], *) bağlantıların gösterdiği öğeler
    hedef = {}
    for ad, soup in belgeler:
        for a in soup.find_all("a", href=True):
            yazi = re.sub(r"\s+", "", a.get_text(strip=True))  # "( 2 )" -> "(2)"
            tip = (a.get("epub:type") or "") + " " + (a.get("role") or "")
            if "#" in a["href"] and ("noteref" in tip or re.fullmatch(r"[\[\(]?\d{1,3}[\]\)]?|\*{1,3}|[¹²³⁴-⁹⁰]+", yazi)):
                dosya, frag = a["href"].split("#", 1)
                hedef[(os.path.basename(dosya) or ad, frag)] = None
    notlar, not_elemanlari, gid_of = {}, set(), {}
    for ad, soup in belgeler:
        for el in soup.find_all(id=True):
            if (ad, el["id"]) in hedef and not (el.name == "a" and el.has_attr("href")):
                kap = el
                if el.name in ("a", "span", "sup"):  # hedef bağlantının kendisi: notun paragrafını al
                    kap = el.find_parent(["p", "li", "aside", "div", "dd"]) or el
                metin = kap.get_text(" ", strip=True)
                metin = re.sub(r"^[\[\(]?\d{1,3}[\]\)]?[.\s]*", "", metin).strip()
                gid = "n%04d" % (len(notlar) + 1)
                notlar[gid] = metin
                gid_of[(ad, el["id"])] = gid
                not_elemanlari.add(id(kap))
        for el in soup.find_all(["aside", "section", "div", "ol"]):
            tip = (el.get("epub:type") or "") + " " + (el.get("role") or "")
            if re.search(r"footnote|endnote|rearnote", tip):
                not_elemanlari.add(id(el))

    ogeler = []
    fihrist = _epub_fihristi(book, belgeler)  # kitabın kendi fihristi: id(blok) -> (seviye, başlık, önüne_ekle)

    def metin_cikar(el, ad):
        out = []
        for c in el.descendants:
            if isinstance(c, NavigableString):
                if c.parent and c.parent.name in ("script", "style"):
                    continue
                if any(p.name == "a" and p.get("href", "").count("#") and
                       (os.path.basename(p["href"].split("#")[0]) or ad, p["href"].split("#", 1)[1]) in gid_of
                       for p in c.parents if p is not el and p.name == "a"):
                    continue
                out.append(str(c))
            elif c.name == "a" and "#" in c.get("href", ""):
                dosya, frag = c["href"].split("#", 1)
                g = gid_of.get((os.path.basename(dosya) or ad, frag))
                if g:
                    out.append("{{" + g + "}}")
            elif c.name in ("span", "a", "div") and re.search(r"pagebreak", (c.get("epub:type") or "") + (c.get("role") or "")):
                et = c.get("aria-label") or c.get("title") or c.get_text(strip=True) or re.sub(r"\D", "", c.get("id", ""))
                if et:
                    out.append("\ue002" + _rakam(et.strip()) + "\ue003")
            elif c.name == "br":
                out.append(" ")
        return TS.norm("".join(out))

    ilk_belge = fihrist["ilk_belge"] if fihrist else 0
    toc_satirlari = []  # baştaki basılı içindekiler sayfası: kitabın sonuna
    for ad, soup in belgeler[:ilk_belge]:
        govde = soup.body or soup
        bl = [b for b in govde.find_all(_BLOK) if not b.find(_BLOK) and b.get_text(strip=True)]
        if bl and (any(TS.ICINDEKILER.search(b.get_text(" ", strip=True)) for b in bl[:3]) or
                   sum(1 for b in bl if b.find("a", href=True)) >= 0.5 * len(bl)):
            toc_satirlari += [TS.norm(b.get_text(" ", strip=True)) for b in bl
                              if not _icindekiler_basligi(b.get_text(" ", strip=True))]
    son_seviye = 1  # fihristteki son başlığın seviyesi (fihristte olmayan gerçek başlık etiketi bunun altına)
    for sira, (ad, soup) in enumerate(belgeler):
        if sira < ilk_belge:
            continue  # fihristin gösterdiği ilk yerden önceki kapak, künye, içindekiler sayfası
        govde = soup.body or soup
        for el in govde.find_all(_BLOK):
            if el.find(_BLOK) or any(id(p) in not_elemanlari for p in [el] + list(el.parents)):
                continue
            metin = metin_cikar(el, ad)
            hedef_f = fihrist["bloklar"].get(id(el)) if fihrist else None
            if hedef_f:
                seviye, baslik, onune = hedef_f
                son_seviye = seviye
                if onune:  # gösterilen yer sıradan paragraf: başlık önüne eklenir, paragraf kaybolmaz
                    ogeler.append({"tur": "baslik", "metin": baslik, "boy": 0, "seviye": seviye, "fihrist": True})
                else:     # başlığın kendisi: metni fihristteki temiz yazı (sayfa işaretleri korunur)
                    isaret = "".join("\ue002" + e + "\ue003" for e in SAYFA_ISARET.findall(metin))
                    ogeler.append({"tur": "baslik", "metin": isaret + baslik, "boy": 0, "seviye": seviye, "fihrist": True})
                    continue
            if not metin or (TS.PAGE_NUM.match(metin) and "\ue002" not in metin):
                continue
            if el.name[0] == "h" and el.name[1:].isdigit() and not fihrist:
                ogeler.append({"tur": "baslik", "metin": metin, "boy": 7 - int(el.name[1])})
            elif el.name[0] == "h" and el.name[1:].isdigit() and not _NOT_BASLIK.match(metin):
                # fihristte olmayan gerçek başlık etiketi (h1-h6): fihristteki son başlığın bir alt seviyesi
                ogeler.append({"tur": "baslik", "metin": metin, "boy": 0, "seviye": min(son_seviye + 1, 3), "fihrist": True})
            else:
                ogeler.append({"tur": "p", "metin": metin, "boy": 0})
    if not fihrist:  # kitabın fihristi yoksa: Stüdyo'nun kuralı (ilk içerik başlığından başlat) ve ön temizlik
        yazilar = [SAYFA_ISARET.sub("", o["metin"]) for o in ogeler]
        temiz = TS.paragraflari_bastan_temizle(yazilar)
        ogeler = _on_temizlik(ogeler[len(yazilar) - len(temiz):])
    else:
        ogeler = [o for i, o in enumerate(ogeler) if not (o["tur"] == "baslik" and not o.get("fihrist") and
                                                         _NOT_BASLIK.match(o["metin"]))]
    # EPUB'ın içindeki basılı içindekiler/fihrist bölümü kitaptan silinir (0.5.8)
    bilgi = {"baslik": (book.get_metadata("DC", "title") or [[""]])[0][0],
             "yazar": (book.get_metadata("DC", "creator") or [[""]])[0][0], "ocr": 0, "bozuk_katman": False,
             "yapi": "fihrist" if fihrist else None, "kapak_resmi": _epub_kapak(book)}
    return ogeler, notlar, bilgi


def _epub_kapak(book):
    """EPUB'un kendi kapak resmi: (bayt, tür) ya da None."""
    from ebooklib import ITEM_COVER, ITEM_IMAGE
    kapak_id = None
    for _, ozellik in book.get_metadata("OPF", "cover") or []:
        kapak_id = (ozellik or {}).get("content") or kapak_id
    adaylar = []
    for it in book.get_items():
        tur = getattr(it, "media_type", "") or ""
        if not tur.startswith("image/"):
            continue
        ad = (it.get_id() or "") + " " + (it.get_name() or "")
        if it.get_type() == ITEM_COVER or it.get_id() == kapak_id:
            return (it.get_content(), tur)
        if "cover" in ad.lower() or "kapak" in ad.lower():
            adaylar.append(it)
    return (adaylar[0].get_content(), adaylar[0].media_type) if adaylar else None


def _sikistir(t):
    """Karşılaştırma için: şapkasız, boşluksuz, küçük harf ('İ SLÂMIN ... R İ' = 'İSLÂMIN ... Rİ')."""
    from .katalog import sade
    return sade(t).replace(" ", "")


def _benzer_baslik(metin, baslik):
    """Gösterilen blok başlığın kendisi mi? Aynı kelimelerle başlasa bile başlıktan çok uzun paragraf başlık değildir."""
    a, b = _sikistir(SAYFA_ISARET.sub("", metin)), _sikistir(baslik)
    if not (a and b) or len(a) > len(b) * 1.3 + 12:
        return False
    return a == b or a.startswith(b) or b.startswith(a) or (len(b) > 8 and b in a)


def _epub_fihristi(book, belgeler):
    """EPUB'un kendi fihristi (nav ya da NCX): her satırın gösterdiği blok başlık olur, girintisi seviyesi.
    Döndürür {'bloklar': {id(blok): (seviye, başlık, önüne_ekle)}, 'ilk_belge': sıra} ya da None (fihrist yok/yetersiz)."""
    duz = []

    def gez(ogeler, derinlik):
        for o in ogeler:
            if isinstance(o, tuple):
                bas, alt = o
                if getattr(bas, "href", None):
                    duz.append((derinlik, bas.title, bas.href))
                gez(alt, derinlik + 1)
            elif getattr(o, "href", None):
                duz.append((derinlik, o.title, o.href))
    gez(book.toc or [], 0)
    duz = [(d, TS.norm(t or ""), h) for d, t, h in duz if (t or "").strip()]
    if len(duz) < 2:
        return None
    sira_of = {ad: i for i, (ad, _) in enumerate(belgeler)}
    bloklar, ilk = {}, None
    for derinlik, baslik, href in duz:
        dosya, _, frag = href.partition("#")
        dosya = os.path.basename(dosya)
        if dosya not in sira_of:
            continue
        soup = belgeler[sira_of[dosya]][1]
        govde = soup.body or soup
        if frag:
            el = govde.find(id=frag)
            if el is None:
                continue
            blok = el if el.name in _BLOK else (el.find_parent(_BLOK) or el.find_next(_BLOK))
        else:  # dosyanın başı: ilk dolu blok
            blok = next((b for b in govde.find_all(_BLOK) if not b.find(_BLOK) and b.get_text(strip=True)), None)
        if blok is None or id(blok) in bloklar:
            continue
        onune = not _benzer_baslik(blok.get_text(" ", strip=True), baslik)
        bloklar[id(blok)] = (min(derinlik + 1, 3), baslik, onune)
        ilk = sira_of[dosya] if ilk is None else min(ilk, sira_of[dosya])
    if len(bloklar) < 2:
        return None
    return {"bloklar": bloklar, "ilk_belge": ilk or 0}


def docx_oku(yol, ilerleme=None):
    _gerekli()
    import docx
    d = docx.Document(yol)
    ogeler, stil_baslik = [], ""
    for p in d.paragraphs:
        t = TS.norm(p.text)
        if not t or TS.PAGE_NUM.match(t):
            continue
        stil = (p.style.name or "").lower() if p.style is not None else ""
        m = re.search(r"(heading|başlık|baslik)\s*(\d)", stil)
        if stil in ("title", "konu başlığı", "kitap adı"):
            stil_baslik = stil_baslik or t  # kitabın adı künyeye gider, metne değil
        elif m:
            ogeler.append({"tur": "baslik", "metin": t, "boy": 7 - int(m.group(2))})
        else:
            ogeler.append({"tur": "p", "metin": t, "boy": 0})
    yazilar = [o["metin"] for o in ogeler]
    temiz = TS.paragraflari_bastan_temizle(yazilar)
    cp = d.core_properties
    return _on_temizlik(ogeler[len(yazilar) - len(temiz):]), {}, {"baslik": cp.title or stil_baslik, "yazar": cp.author or "",
                                                                  "ocr": 0, "bozuk_katman": False}


def txt_oku(yol, ilerleme=None):
    _gerekli()
    paras = TS.from_txt_bytes(open(yol, "rb").read())
    ogeler = []
    for t in TS.paragraflari_bastan_temizle(paras):  # noqa: B007
        harf = [c for c in t if c.isalpha()]
        b = len(t) <= 60 and (TS.ICERIK_BASLIK.match(t) or BOLUM_NO.match(t) or (harf and all(c.isupper() for c in harf)))
        ogeler.append({"tur": "baslik" if b else "p", "metin": t, "boy": 0})
    return _on_temizlik(ogeler), {}, {"baslik": "", "yazar": "", "ocr": 0, "bozuk_katman": False}


# ======================= ortak: düzeltme ve kitap.json =======================
def _buyuk_harfli(t):
    h = [c for c in SAYFA_ISARET.sub("", t) if c.isalpha()]
    return bool(h) and all(c.isupper() for c in h)


_COP_ISARET = re.compile(r"[#»«|<>@^~_=\\{}\[\]]")


def anlamli_baslik(t):
    """Başlık gerçekten başlık mı? OCR çöpü (Arapça satırın Türkçe OCR'ı) ve tek kalmış cümle sonları elenir."""
    t = SAYFA_ISARET.sub("", t).strip()
    isaretli = bool(re.search(r"(?:\ue000\d{1,3}\ue001|\{\{n\d+\}\})\s*$", t))
    t = re.sub(r"(?:\ue000\d{1,3}\ue001|\{\{n\d+\}\})\s*$", "", t).strip()  # 0.5.15: "İLİM²"
    if not t or len(t) > 90 or _COP_ISARET.search(t) or t[0] in "“\"'‘«(":
        return False  # tırnakla başlayan satır cümle parçasıdır
    if TS.ICERIK_BASLIK.match(t) or BOLUM_NO.match(t) or _harfli_madde_basligi(t) or re.fullmatch(
            r"(bölüm|kısım|fasıl|bab|kitap|makale)\s+([ivxlc]{1,6}|\d{1,3})[.:]?", t, re.I):
        return True
    if not t[0].isalnum():
        return False  # ? ile başlayan vb.
    if len(t.split()) == 1 and len(re.sub(r"[^\w]", "", t)) < 5 and not (
            isaretli and t.isupper() and len(re.sub(r"[^\w]", "", t)) >= 3 and _gecerli_kelime(t)):
        # (dipnot işaretli kısa büyük harfli kelime başlıktır: "İLİM²"; işaretsiz "TİRE" gibi kırıntı değil)
        return False  # tek kelimelik kısa satır (TİRE): bilinen başlıklar yukarıda kabul edildi
    if t[0].islower() or (t.endswith((".", ",", ";")) and not re.search(r"\b(vs|bkz|s|c)\.$", t, re.I)):
        return False
    kelimeler = re.findall(r"[^\W\d_]{2,}", t)
    if not kelimeler or max(len(k) for k in kelimeler) < 3:
        return False
    harf = sum(len(k) for k in kelimeler)
    if harf / max(1, len(re.sub(r"\s", "", t))) < 0.7:
        return False
    baglac = ("ve", "ile", "ki", "da", "de", "ya", "veya", "ya da")
    asil_kel = [k for k in kelimeler if k.lower() not in baglac] or kelimeler
    if sum(len(k) for k in asil_kel) / len(asil_kel) < 3.5:
        return False  # OCR çöpü: 2-3 harflik parçalar
    for parca in t.split():  # tek başına duran harf (EĞEN ğ): çöp; noktalı kısaltmalar (S.A.V.) muaf
        oz = parca.strip("()[]:;,!?\"“”'’‘")
        if len(oz) == 1 and oz.isalpha() and oz.lower() != "o" and not re.fullmatch(r"[IVXLC]", oz):
            return False
    # büyük/küçük harf düzeni: TAMAMI BÜYÜK ya da Düzgün Yazım; karışık (PAS TAİ Kan yay) başlık değildir
    buyuk = [k for k in kelimeler if len(k) >= 2 and k.isupper()]
    kucuk = [k for k in kelimeler if not k.isupper() and k.lower() not in ("ve", "ile", "ki", "da", "de", "ya", "veya")]
    if buyuk and kucuk and not all(re.search(re.escape(k) + r"\.", t) for k in buyuk):
        return False
    iyi = sum(1 for k in kelimeler if DZ.gecerli_mi(k) or DZ._kelime_mi(k) or
              (k.isupper() and _turkce_kelime_onar(k) != k))  # 0.5.15: "BESINCI" (BEŞİNCİ)
    if len(kelimeler) >= 2 and buyuk and not kucuk:  # TAMAMI BÜYÜK: dini terimler sözlükte yok (RUBÛBİYYETİ BİLMEK)
        return iyi / len(kelimeler) >= 0.5
    return iyi / len(kelimeler) >= 0.75


def _basliklari_denetle(ogeler):
    """Anlamsız başlıklar paragraf olur; küçük harfle başlayan cümle sonu önceki paragrafa eklenir."""
    out = []
    for k, o in enumerate(ogeler):
        sonraki = ogeler[k + 1] if k + 1 < len(ogeler) else None
        if o.get("fihrist") or o.get("koru"):  # kitabın kendi fihristinden / basılı içindekiler: olduğu gibi
            out.append(o)
            continue
        if o["tur"] == "baslik" and sonraki and sonraki["tur"] == "p" and \
                SAYFA_ISARET.sub("", sonraki["metin"]).lstrip()[:1].islower():
            # arkasından küçük harfle devam eden paragraf: bu satır cümlenin başıdır, başlık değil
            sonraki["metin"] = o["metin"] + " " + sonraki["metin"]
            continue
        if o["tur"] == "baslik" and not anlamli_baslik(o["metin"]):
            o = dict(o, tur="p")
            yazi = SAYFA_ISARET.sub("", o["metin"]).strip()
            onceki = out[-1] if out else None
            if onceki and onceki["tur"] == "p" and yazi[:1].islower() and not onceki["metin"].endswith(END_PUNCT):
                onceki["metin"] += " " + o["metin"]
                continue
        out.append(o)
    return out


_EK_BASLIK = re.compile(r"^(sonuç|hâtime|hatime|netice|takdim|dîbâce|dibace|kaynakça|bibliyografya|sözlük|lügatçe|"
                        r"dizin|indeks|ekler?|mukaddime|önsöz|giriş|başlangıç)[.:]?$", re.I)


def _kalip_baslik(t):
    t = SAYFA_ISARET.sub("", t).strip()
    return bool(TS.ICERIK_BASLIK.match(t) or BOLUM_NO.match(t) or _EK_BASLIK.match(t) or re.match(
        r"(bölüm|kısım|fasıl|bab|kitap|makale)\s+([ivxlc]{1,6}|\d{1,3})\b", t, re.I))


def _seviyeler(ogeler):
    _seviyeler_ic(ogeler)
    for o in ogeler:
        if o.get("koru") and o["tur"] == "baslik":
            o["seviye"] = 1


def _seviyeler_ic(ogeler):
    if any(o.get("fihrist") for o in ogeler):
        for o in ogeler:  # kitabın kendi fihristi: seviyeler oradan; fihristte olmayan başlık paragraf olur
            if o["tur"] == "baslik" and not o.get("fihrist") and not o.get("koru"):
                o["tur"] = "p"
        return
    _seviyeler_boy(ogeler)
    basliklar = [o for o in ogeler if o["tur"] == "baslik"]
    kalip = [o for o in basliklar if _kalip_baslik(o["metin"])]
    if len(kalip) >= 2:  # "Birinci Bölüm", "Önsöz"... her zaman en üstte; öteki başlıklar altında (en az 2. seviye)
        for o in basliklar:
            o["seviye"] = 1 if o in kalip else max(2, o["seviye"])


def _seviyeler_boy(ogeler):
    """Başlık seviyesi (en çok 3). Metin katmanında punto güvenilir: boylardan. OCR'da satır yüksekliği harflere göre
    oynar: önce yazım biçimi (TAMAMI BÜYÜK HARF üst seviye), aynı biçim içinde %20'den büyük punto farkı."""
    basliklar = [o for o in ogeler if o["tur"] == "baslik"]
    if basliklar and sum(1 for o in basliklar if o.get("ocr")) > len(basliklar) / 2:
        gruplar = []  # (büyük_harf_mı, boy eşiği)
        for buyuk in (True, False):
            boylar = sorted({o["boy"] for o in basliklar if _buyuk_harfli(o["metin"]) == buyuk and o["boy"]}, reverse=True)
            esik = []
            for b in boylar:
                if not esik or b < esik[-1] * 0.8:
                    esik.append(b)
            gruplar += [(buyuk, e) for e in esik]
        for o in basliklar:
            buyuk = _buyuk_harfli(o["metin"])
            aday = [k for k, (g, e) in enumerate(gruplar) if g == buyuk and o["boy"] >= e * 0.8]
            o["seviye"] = min((aday[0] if aday else len(gruplar) - 1) + 1, 3)
        # boşluk kalmasın: kullanılan seviyeler 1, 2, 3 diye sıkıştırılır
        kullanilan = sorted({o["seviye"] for o in basliklar})
        for o in basliklar:
            o["seviye"] = kullanilan.index(o["seviye"]) + 1
        return
    """Başlık boylarından fihrist seviyesi (en büyük = 1, en çok 3). Boy bilgisi yoksa 1."""
    boylar = sorted({o["boy"] for o in ogeler if o["tur"] == "baslik" and o["boy"]}, reverse=True)
    esikler = []  # her seviyenin en büyük boyu; %12'den fazla küçülünce yeni seviye
    for b in boylar:
        if not esikler or b < esikler[-1] * 0.88:
            esikler.append(b)
    for o in ogeler:
        if o["tur"] == "baslik":
            if o["boy"] and esikler:
                o["seviye"] = min(next((k for k, e in enumerate(esikler) if o["boy"] >= e * 0.88), len(esikler) - 1) + 1, 3)
            else:
                o["seviye"] = 1


def _basliklari_birlestir(ogeler):
    """'BİRİNCİ BÖLÜM' + hemen ardındaki aynı seviyeli başlık -> 'BİRİNCİ BÖLÜM: Başlık'."""
    out = []
    for o in ogeler:
        p = out[-1] if out else None
        if p and p["tur"] == o["tur"] == "baslik" and p["seviye"] == o["seviye"] and not p.get("fihrist") and \
                BOLUM_NO.match(SAYFA_ISARET.sub("", p["metin"]).strip()) and not BOLUM_NO.match(o["metin"]):
            p["metin"] = p["metin"].rstrip(" .:") + ": " + SAYFA_ISARET.sub("", o["metin"]).strip()
            p["metin"] += "".join("\ue002" + e + "\ue003" for e in SAYFA_ISARET.findall(o["metin"]))
            continue
        out.append(o)
    return out


def _duzelt(metinler):
    """Stüdyo düzeltmeleri (sayfa atıfları KORUNUR: kitabın basılı sayfa numaraları EPUB'da var)."""
    basliklar = DZ.satir_ici_ust_bilgileri_bul([SAYFA_ISARET.sub(" ", t) for t in metinler])
    out = []
    for p in metinler:
        p = DZ.satir_ici_ust_bilgileri_sil(p, basliklar)
        p = DZ.cop_isaretleri_sil(p)
        p = re.sub(r"(?<=[^\W\d_])-\s*(\ue002[^\ue003]{1,20}\ue003)\s*(?=[a-zçğıöşüâîû])", r"\1", p)  # 0.5.13
        p = DZ.satir_ici_tireleri_birlestir(p)
        if hasattr(DZ, "bolunmus_kelimeleri_birlestir"):  # "oldu ğundan" -> "olduğundan" (Stüdyo'nun onarımı)
            p = DZ.bolunmus_kelimeleri_birlestir(p)
        p = re.sub(r"(?<=[^\W\d_])\s+([’'])\s*(?=[^\W\d_])", r"\1", p)  # Nasır ’ ın -> Nasır’ın
        p = re.sub(r"\(\s*(\{\{n\d{4,}\}\})\s*\)", r"\1", p)         # ( {{n0002}} ) -> {{n0002}}
        p = re.sub(r"\s*\(\s*\)", "", p)                                  # silinen numaradan kalan "()"
        p = DZ.harfleri_onar(p)
        p = re.sub(r"[ \t]{2,}", " ", p).replace(" ,", ",").replace(" .", ".").strip()
        p = re.sub(r"^(\d{1,3}[.)])(?=[^\s\d.)\ue002])", r"\1 ", p)
        out.append(p)
    return out


def _konumlar(metin):
    """Sayfa işaretlerini metinden çıkarır: (temiz metin, [(etiket, konum)])."""
    sayfalar, parcalar, uz = [], [], 0
    for i, parca in enumerate(SAYFA_ISARET.split(metin)):
        if i % 2:
            sayfalar.append((parca, uz))
        else:
            parcalar.append(parca)
            uz += len(parca)
    duz = "".join(parcalar)
    # baştaki/sondaki boşluk kırpılınca konumlar kayar: düzelt
    sol = len(duz) - len(duz.lstrip())
    duz = duz.strip()
    return duz, [(e, max(0, min(len(duz), k - sol))) for e, k in sayfalar]


def kitaba_cevir(ogeler, notlar, kunye):
    ogeler = _basliklari_denetle(ogeler)
    _seviyeler(ogeler)
    # 0.5.14: "A — ZAMAN", "B — İRTİKÂ", "C - RIZIK": harfle numaralı ara başlıklar aynı seviyede (yazı boyu tahmini
    # birini bir alt seviyeye düşürebiliyor); en sık seviye, eşitlikte üstteki
    harfli = [o for o in ogeler if o["tur"] == "baslik" and not o.get("fihrist") and _harfli_madde_basligi(o["metin"])]
    if len(harfli) >= 2:
        sayim = collections.Counter(o.get("seviye", 1) for o in harfli).most_common()
        ortak = min(sv for sv, c in sayim if c == sayim[0][1])
        for o in harfli:
            o["seviye"] = ortak
    ogeler = _basliklari_birlestir(ogeler)
    ogeler = [o for o in ogeler if SAYFA_ISARET.sub("", o["metin"]).strip() or SAYFA_ISARET.search(o["metin"])]
    # tek başına kalmış ayet/madde numarası: sonraki paragrafa (Stüdyo kuralı)
    birlesik, bekleyen = [], None
    for o in ogeler:
        yazi = SAYFA_ISARET.sub("", o["metin"]).strip()
        if o["tur"] == "p" and re.fullmatch(r"\d{1,3}\s*[.)]", yazi):
            bekleyen = o["metin"]
            continue
        if bekleyen:
            o = dict(o, metin=bekleyen.strip() + " " + o["metin"])
            bekleyen = None
        birlesik.append(o)
    ogeler = birlesik
    duz = _duzelt([o["metin"] for o in ogeler])
    notlar = dict(zip(notlar, _duzelt(list(notlar.values())))) if notlar else {}
    kit = K.yeni(kunye)
    tasinan = []  # atılan paragraftaki sayfa işaretleri sonrakine geçer
    for o, metin in zip(ogeler, duz):
        temiz, sayfalar = _konumlar(metin)
        yazi = K.NOT_ISARETI.sub("", temiz).strip()
        if o["tur"] == "p" and not o.get("koru") and not K.NOT_ISARETI.search(temiz) and (not yazi or (DZ.cop_paragraf_mi(yazi) and not _AR.search(yazi))):
            tasinan += [e for e, _ in sayfalar]
            continue
        sayfa = [{"no": e, "konum": {"tr": 0}} for e in tasinan] + [{"no": e, "konum": {"tr": k}} for e, k in sayfalar]
        tasinan = []
        if o["tur"] == "baslik":
            temiz2 = temiz if o.get("fihrist") and not o.get("onar") else turkce_onar(_buyuk_baslik_onar(temiz))  # kitabın kendi fihristindeki yazı zaten temiz
            if len(temiz2) == len(temiz):  # uzunluk aynı kalır (harf değişimi): sayfa konumları geçerli
                temiz = temiz2
            K.blok_ekle(kit, "baslik", {"tr": temiz}, seviye=o.get("seviye", 1), sayfalar=sayfa)
        else:
            K.blok_ekle(kit, "p", {"tr": temiz}, sayfalar=sayfa)
    kullanilan = set()
    for b in kit["bloklar"]:
        kullanilan |= set(K.NOT_ISARETI.findall(b["metin"]["tr"]))
    kit["dipnotlar"] = {g: {"metin": {"tr": t}} for g, t in notlar.items() if g in kullanilan}
    # aynı etiketin tekrarı (ör. boş sayfa) ve sıra bozukluğu: ilk görüleni tut
    gorulen = set()
    for b in kit["bloklar"]:
        if "sayfalar" in b:
            b["sayfalar"] = [s for s in b["sayfalar"] if not (s["no"] in gorulen or gorulen.add(s["no"]))]
            if not b["sayfalar"]:
                del b["sayfalar"]
    _sirayi_duzelt(kit)
    return kit


def _sirayi_duzelt(kit):
    """Sayfa numaraları geriye gidiyorsa (yanlış okunmuş numara) o işaretler atılır."""
    onceki = None
    for b in kit["bloklar"]:
        if "sayfalar" not in b:
            continue
        tut = []
        for s in b["sayfalar"]:
            if s["no"].isdigit():
                if onceki is not None and int(s["no"]) <= onceki:
                    continue
                onceki = int(s["no"])
            tut.append(s)
        b["sayfalar"] = tut
        if not tut:
            del b["sayfalar"]


def _kapak_basligi(ilk_sayfalar):
    """İlk sayfalardaki en büyük puntolu yazı (kapak başlığı)."""
    dolu = next((rows for rows in ilk_sayfalar if any(_harf(r["text"]) >= 3 for r in rows)), [])
    aday = [r for r in dolu if 3 <= len(r["text"]) <= 80 and _harf(r["text"]) >= 3]
    if not aday:
        return ""
    en = max(r["h"] for r in aday)
    return " ".join(r["text"] for r in aday if r["h"] >= en * 0.95)[:120]


_KUCUK_KAL = {"ve", "ile", "ki", "veya", "ya", "da", "de", "fi", "min", "an", "ala", "li"}


def turkce_baslik(s):
    """'İTİKADDA ORTA YOL' -> 'İtikadda Orta Yol' (Türkçe büyük/küçük harf kurallarıyla)."""
    harf = [c for c in s if c.isalpha()]
    if not harf or not all(c.isupper() for c in harf):
        return s
    kucuk = s.replace("I", "ı").replace("İ", "i").lower()
    out = []
    for i, w in enumerate(kucuk.split()):
        if i and w in _KUCUK_KAL:
            out.append(w)
        else:
            out.append({"i": "İ", "ı": "I"}.get(w[0], w[0].upper()) + w[1:])
    return " ".join(out)


def _benzer(a, b):
    from .katalog import sade
    ka, kb = set(sade(a).split()), set(sade(b).split())
    return bool(ka and kb) and len(ka & kb) / min(len(ka), len(kb)) >= 0.5


_PROGRAM_ADI = re.compile(
    r"design|adobe|acrobat|microsoft|office|word|writer|windows|abbyy|finereader|scan|tarayıcı|epson|canon|\bhp\b|"
    r"printer|bullzip|pdf|corel|quark|indesign|calibre|\buser\b|admin|owner|kullanıcı|bilgisayar|\bpc\b|www|\.com|"
    r"unknown|bilinmiyor|anonymous|construction|default|untitled", re.I)


def _kisi_adi_mi(t):
    """PDF/EPUB bilgi alanındaki yazar bir insan adına benziyor mu ("Construction.design", "Microsoft Word - …",
    "Administrator" gibi program/bilgisayar adları değil)."""
    t = (t or "").strip()
    return bool(t) and not re.search(r"[.@/\\\d_]|https?:", t.replace(". ", " ").rstrip(".")) and \
        len(t.split()) <= 6 and not _PROGRAM_ADI.search(t)


def _kunye_sec(bilgi, ad_baslik, ad_yazar, dosya_koku):
    """Eser adı: kapaktaki başlık (Türkçe harfleriyle; dosya adıyla uyuşuyorsa) > PDF/EPUB bilgi alanı (dosya adının
    kopyası değilse) > dosya adı. Yazar: bilgi alanı (eser adıyla aynı değilse) > dosya adındaki 'Yazar - Eser'."""
    from .katalog import sade
    # bilgi alanı ancak dosya adındaki "Yazar - Eser"in birleşik kopyasıysa atılır; sadece eser adıysa (çoğu zaman
    # Türkçe harfleriyle daha doğru yazılmıştır) tercih edilir
    kopya = lambda s: bool(ad_yazar) and sade(s) == sade(ad_yazar + " " + ad_baslik)
    meta_b = bilgi.get("baslik") if _anlamli(bilgi.get("baslik")) else ""
    # 0.5.15: bilgi alanındaki ad dosya adındaki eser adıyla hiç kelime paylaşmıyorsa çöp ya da başka bir eserdir
    ad_kel = {k for k in sade(ad_baslik or "").split() if len(k) >= 3}
    if meta_b and len(ad_kel) >= 2 and not ad_kel & {k for k in sade(meta_b).split() if len(k) >= 3}:
        meta_b = ""
    meta_y = bilgi.get("yazar") if _anlamli(bilgi.get("yazar")) and _kisi_adi_mi(bilgi.get("yazar")) else ""
    satirlar = bilgi.get("kapak_satirlari") or ([(bilgi["kapak_baslik"], 1, 0)] if bilgi.get("kapak_baslik") else [])
    if not ad_baslik.isascii():  # dosya adında Türkçe harfler var: kullanıcının verdiği düzgün ad, en güvenilir kaynak
        ek = r"(nin|nın|nun|nün|in|ın|un|ün|a|e|ya|ye|da|de|ta|te|dan|den|tan|ten|la|le|yla|yle|ı|i|u|ü|yı|yi|yu|yü)"
        duzgun = lambda t: re.sub(r"(?<=[^\W\d_])-(?=" + ek + r"\b)", "’", re.sub(r"\s+([)\]])", r"\1", t)).strip()
        yazar = meta_y if meta_y and not kopya(meta_y) else turkcelestir(ad_yazar)
        return duzgun(ad_baslik), duzgun(yazar)
    kapak = _kapak_sec(satirlar, ad_baslik if ad_yazar else "")  # dosya adı "Yazar - Eser" değilse karşılaştırılamaz
    if kapak:
        baslik = turkce_baslik(kapak)
    elif meta_b and not kopya(meta_b):
        baslik = meta_b
    else:
        baslik = turkcelestir(ad_baslik)
    if meta_y and not kopya(meta_y) and sade(meta_y) != sade(baslik) and sade(baslik) not in sade(meta_y):
        yazar = meta_y
    else:
        yazar = turkcelestir(ad_yazar)
    return turkce_onar(baslik), turkce_onar(yazar)


def _kapak_sec(satirlar, ad_baslik):
    """Kapaktaki başlık satırları: en büyük satırdan başlayıp alttaki (ya da üstteki) büyük satırlar eklenir; dosya
    adındaki kelimelerin en az %80'ini içeren en kısa birleşim seçilir. Dosya adı yoksa en büyük satır(lar)."""
    from .katalog import sade
    if not satirlar:
        return ""
    en = max(h for _, h, _ in satirlar)
    buyuk = [(t.strip(" ,;:"), h) for t, h, _ in satirlar if h >= en * 0.5]
    hedef = set(sade(ad_baslik).split())
    if not hedef:
        return " ".join(t for t, h in buyuk if h >= en * 0.95)
    iyi, iyi_puan = "", 0
    for bas in range(len(buyuk)):
        for son in range(bas + 1, min(len(buyuk), bas + 4) + 1):
            metin = " ".join(t for t, _ in buyuk[bas:son])
            kel = set(sade(metin).split())
            kapsama = len(kel & hedef) / len(hedef)
            puan = kapsama - 0.05 * len(kel - hedef)
            if kapsama >= 0.8 and puan > iyi_puan:
                iyi, iyi_puan = metin, puan
    return iyi


_ESLER = {"o": "ö", "ö": "o", "u": "ü", "ü": "u", "s": "ş", "ş": "s", "c": "ç", "ç": "c", "g": "ğ", "ğ": "g", "i": "ı", "ı": "i"}


def _tr_kucuk(w):
    return w.replace("I", "ı").replace("İ", "i").lower()


def _tr_buyuk(w):
    return w.replace("i", "İ").replace("ı", "I").upper()


def _gecerli(w):
    kelimeler, _ = DZ._sozluk()
    return w in kelimeler or DZ.gecerli_mi(w)


def _kelime_onar(kucuk):
    """Geçersiz kelimenin noktası/şapkası kaybolmuş doğru hâli: önce kelime listesindeki iskelet eşi, yoksa Zemberek'e
    sorarak (en az değişiklikle). Bulunamazsa None."""
    if len(kucuk) < 2 or _gecerli(kucuk):
        return None
    _, iskelet = DZ._sozluk()
    aday = iskelet.get(kucuk.translate(DZ._TR_ISKELET))
    if aday and aday != kucuk:
        return aday
    yerler = [i for i, c in enumerate(kucuk) if c in _ESLER][:6]
    import itertools
    for adet in range(1, len(yerler) + 1):
        for secim in itertools.combinations(yerler, adet):
            k = list(kucuk)
            for i in secim:
                k[i] = _ESLER[k[i]]
            k = "".join(k)
            if DZ.gecerli_mi(k):
                return k
    return None


def turkce_onar(metin):
    """Başlık, kitap adı ve yazar için harf onarımı (OCR'ın kaybettiği nokta/şapka: 'SÖZÜN OZÜ' -> 'SÖZÜN ÖZÜ').
    Büyük/küçük harf düzeni korunur; karşılığı bulunamayan kelimeye (özel ad) dokunulmaz."""
    if not metin or DZ is None:
        return metin

    def onar(m):
        w = m.group(0)
        if re.fullmatch(r"[IVXLCDM]+", w):  # Roma rakamı (II, IV) kelime değildir
            return w
        dogru = _kelime_onar(_tr_kucuk(w))
        if not dogru:
            return w
        if w.isupper():
            return _tr_buyuk(dogru)
        if w[:1].isupper():
            return _tr_buyuk(dogru[0]) + dogru[1:]
        return dogru
    return re.sub(r"[^\W\d_]+", onar, metin)


def turkcelestir(s):
    """Şapkasız/Türkçe harfsiz yazımı (dosya adı) kelime listesiyle düzeltir: 'Itikatta Sozun Ozu' -> 'İtikatta Sözün Özü'."""
    if not s or DZ is None:
        return s
    kelimeler, iskelet = DZ._sozluk()
    out = []
    for w in s.split():
        kk = DZ._kucuk(w)
        if not w.isascii() or not kk.isalpha():
            out.append(w)
            continue
        if w[:1] == "I" and (("i" + kk[1:]) in kelimeler or DZ.gecerli_mi("i" + kk[1:])):  # 'Imam' = İmam
            out.append("İ" + w[1:])
            continue
        if kk in kelimeler:
            out.append(w)
            continue
        dogru = iskelet.get(kk.translate(DZ._TR_ISKELET))
        if not dogru:
            out.append(w)
            continue
        if w[:1].isupper():
            dogru = {"i": "İ", "ı": "I"}.get(dogru[0], dogru[0].upper()) + dogru[1:]
        out.append(dogru)
    return " ".join(out)


def _dosya_adindan(yol):
    ad = os.path.splitext(os.path.basename(yol))[0]
    ad = re.sub(r"[_]+", " ", ad).strip()
    if " - " in ad:
        yazar, baslik = ad.split(" - ", 1)
        return baslik.strip(), yazar.strip()
    return ad, ""


def _anlamli(s):
    s = (s or "").strip()
    return s and not re.search(r"microsoft|word|untitled|adsız|\.docx?|\.pdf|^[\d\W]+$", s, re.I) and len(s) < 150


def cevir(yol, ilerleme=None, kaynak_bilgi=None, kapak_yolu=None):
    """Dosya -> kitap.json sözlüğü (Türkçe). kapak_yolu verilirse kitabın kendi kapak görseli oraya (uzantısıyla) yazılır."""
    uzanti = os.path.splitext(yol)[1].lower()
    okuyucu = {".pdf": pdf_oku, ".epub": epub_oku, ".docx": docx_oku, ".txt": txt_oku}.get(uzanti)
    if not okuyucu:
        raise ValueError("Desteklenmeyen dosya türü: " + uzanti)
    ogeler, notlar, bilgi = okuyucu(yol, ilerleme)
    if ilerleme:
        ilerleme("Metin düzeltiliyor")
    ad_baslik, ad_yazar = _dosya_adindan(yol)
    baslik, yazar = _kunye_sec(bilgi, ad_baslik, ad_yazar, os.path.splitext(os.path.basename(yol))[0])
    kunye = {
        "baslik": {"tr": baslik},
        "yazar": {"tr": yazar},
        "asil_dil": "tr",
        "kaynak": {"tur": "dosya", "ad": os.path.basename(yol), **(kaynak_bilgi or {})},
        "sayfa_kaynagi": "basılı baskı" if uzanti == ".pdf" else "",
        "cikarma": {k: bilgi.get(k) for k in ("sayfa", "ocr", "bozuk_katman")},
        "yapi": bilgi.get("yapi"),  # "fihrist": başlıklar kitabın kendi fihristinden (EPUB) ya da yer imlerinden (PDF)
    }
    kit = kitaba_cevir(ogeler, notlar, kunye)
    if kapak_yolu and bilgi.get("kapak_resmi"):
        veri, tur = bilgi["kapak_resmi"]
        uz = ".png" if "png" in tur else ".jpg"
        with open(kapak_yolu + uz, "wb") as f:
            f.write(veri)
        kit["kunye"]["kapak"] = os.path.basename(kapak_yolu + uz)
    if not any(b["tur"] == "p" for b in kit["bloklar"]):
        raise ValueError("Dosyadan metin çıkarılamadı (boş ya da okunamayan dosya)")
    return kit


def kimlik_uret(yol):
    ad = os.path.splitext(os.path.basename(yol))[0]
    ad = ad.translate(str.maketrans("İıŞşĞğÇçÖöÜü", "IiSsGgCcOoUu"))
    ad = "".join(c for c in unicodedata.normalize("NFKD", ad) if not unicodedata.combining(c))
    ad = re.sub(r"[^A-Za-z0-9]+", "-", ad).strip("-")[:60] or "kitap"
    return f"{ad}-{hashlib.sha1(yol.encode('utf-8')).hexdigest()[:6]}"
